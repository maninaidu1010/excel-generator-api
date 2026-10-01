"""
excel_engine.py - spec-driven Excel service (Flask blueprint) for the Copilot Studio agent
==========================================================================================

Three actions (plus download, which lives in app.py):

  POST /createexcel    build a polished workbook from a declarative spec
  POST /editexcel      open an existing workbook (base64 / URL / CSV), apply a list of
                       operations (format, clean, sort, formulas, charts, validation ...)
  POST /analyzeexcel   inspect + validate a workbook: structure, data quality, formula
                       problems, features that editing could lose
  GET  /excel-capabilities   what the engine can do (for the agent / for you)

Design notes
------------
* One library (openpyxl) for both create and edit, so every feature works in both modes.
* The original file is never modified: edits are saved to a new file.
* Formulas are written as real Excel formulas. openpyxl does not calculate them, so the
  workbook is flagged to fully recalculate when opened in Excel. If LibreOffice happens to
  be installed on the server it is used (read-only copy) to verify formulas; on Render's
  native Python runtime it is not, and the response says formulas were not verified.
* Honest limits are reported back to the agent in `warnings` (e.g. features openpyxl cannot
  preserve: slicers, pivot caches, sparklines, threaded comments, macros without .xlsm).
"""
from __future__ import annotations

import base64
import binascii
import csv
import io
import ipaddress
import math
import os
import re
import shutil
import socket
import statistics
import subprocess
import tempfile
import time
import uuid
import zipfile
from collections import Counter, OrderedDict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple
from urllib.parse import urlparse
import urllib.error
import urllib.request

from flask import Blueprint, jsonify, request
from openpyxl import Workbook, load_workbook
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import (CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, IconSetRule, Rule)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

MAX_FILE_BYTES = int(os.environ.get("EXCEL_MAX_FILE_MB", "15")) * 1024 * 1024
MAX_CELLS_TOUCHED = 400_000
MAX_ROWS_IN_SPEC = 50_000
EXCEL_EXTS = (".xlsx", ".xlsm")


class ExcelError(Exception):
    """User-facing error (safe to show to the agent / end user)."""

    def __init__(self, message: str, status: int = 400, code: str = "EXCEL_BAD_REQUEST", details: Any = None):
        super().__init__(message)
        self.message, self.status, self.code, self.details = message, status, code, details


# --------------------------------------------------------------------------- #
# Themes and number formats
# --------------------------------------------------------------------------- #
THEMES: Dict[str, Dict[str, Any]] = {
    "vkollab": dict(dark="0B3C7D", primary="0F6CBD", light="EAF1FB", band="F5F8FD", border="C9D6EA",
                    text="1B2A41", muted="5B6B8C",
                    accents=["0F6CBD", "0B3C7D", "2B88D8", "3A7CA5", "1B4F9C", "5B9BD5"], font="Calibri"),
    "slate": dict(dark="1F2937", primary="374151", light="F3F4F6", band="F9FAFB", border="D1D5DB",
                  text="111827", muted="6B7280",
                  accents=["374151", "2563EB", "059669", "D97706", "DC2626", "7C3AED"], font="Calibri"),
    "green": dict(dark="14532D", primary="16A34A", light="DCFCE7", band="F0FDF4", border="BBF7D0",
                  text="14281D", muted="4B6355",
                  accents=["16A34A", "14532D", "22C55E", "0D9488", "65A30D", "84CC16"], font="Calibri"),
    "orange": dict(dark="7C2D12", primary="EA580C", light="FFEDD5", band="FFF7ED", border="FED7AA",
                   text="2B1A10", muted="7A5C48",
                   accents=["EA580C", "7C2D12", "F97316", "D97706", "B45309", "FB923C"], font="Calibri"),
}
CURRENCY_SYMBOLS = {"USD": "$", "EUR": "€", "GBP": "£", "INR": "₹", "JPY": "¥", "AUD": "A$", "CAD": "C$"}
HEX6 = re.compile(r"^#?([0-9a-fA-F]{6})$")


def hex6(value: Any, default: Optional[str] = None) -> Optional[str]:
    m = HEX6.match(str(value or "").strip())
    return m.group(1).upper() if m else default


def blend(color: str, ratio: float, with_color: str = "FFFFFF") -> str:
    """Mix `color` toward `with_color` (ratio of with_color)."""
    a, b = [int(color[i:i + 2], 16) for i in (0, 2, 4)], [int(with_color[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x * (1 - ratio) + y * ratio):02X}" for x, y in zip(a, b))


class Theme:
    def __init__(self, spec: Any = None, currency: Optional[str] = None):
        name = "vkollab"
        overrides: Dict[str, Any] = {}
        if isinstance(spec, str):
            name = spec.lower() if spec.lower() in THEMES else "vkollab"
        elif isinstance(spec, dict):
            name = str(spec.get("name", "vkollab")).lower()
            name = name if name in THEMES else "vkollab"
            overrides = spec
            currency = currency or spec.get("currency")
        base = dict(THEMES[name])
        primary = hex6(overrides.get("primary"))
        if primary:  # derive a coherent palette from one brand colour
            base.update(primary=primary, dark=hex6(overrides.get("dark")) or blend(primary, 0.45, "000000"),
                        light=blend(primary, 0.9), band=blend(primary, 0.95), border=blend(primary, 0.75),
                        accents=[primary, blend(primary, 0.45, "000000"), blend(primary, 0.25), blend(primary, 0.5),
                                 blend(primary, 0.3, "000000"), blend(primary, 0.4, "FFFFFF")])
        if overrides.get("font"):
            base["font"] = str(overrides["font"])[:40]
        self.__dict__.update(base)
        self.currency = (currency or "USD").upper()
        self.symbol = CURRENCY_SYMBOLS.get(self.currency, "$")


def number_format(name: Any, theme: Theme) -> Optional[str]:
    if not name:
        return None
    key, s = str(name).strip().lower(), theme.symbol
    presets = {
        "currency": f'"{s}"#,##0.00', "currency0": f'"{s}"#,##0', "integer": "#,##0", "number": "#,##0.00",
        "number1": "#,##0.0", "percent": "0.0%", "percent0": "0%", "percent2": "0.00%",
        "date": "dd-mmm-yyyy", "date_short": "dd/mm/yyyy", "month": "mmm-yyyy", "datetime": "dd-mmm-yyyy hh:mm",
        "time": "hh:mm", "text": "@", "general": "General", "thousands": '#,##0,"K"', "millions": '#,##0.0,,"M"',
        "multiple": '0.0"x"',
        "accounting": f'_("{s}"* #,##0.00_);_("{s}"* (#,##0.00);_("{s}"* "-"??_);_(@_)',
    }
    return presets.get(key, str(name))


CURRENCY_WORDS = ("price", "cost", "revenue", "amount", "sales", "salary", "budget", "spend", "profit", "total",
                  "fee", "income", "expense", "payment", "balance", "value", "turnover", "margin amount")
PERCENT_WORDS = ("%", "percent", "pct", "rate", "margin", "share", "growth", "utilization", "utilisation")
PLAIN_NUMBER_WORDS = ("id", "year", "code", "zip", "pin", "phone", "number", "no.", "sr", "s.no", "qty code")


def is_formula(v: Any) -> bool:
    return isinstance(v, str) and v.startswith("=")


def infer_format(header: str, values: List[Any]) -> Optional[str]:
    """Best-guess number format key for a column from its header and values."""
    h = (header or "").lower().strip()
    vals = [v for v in values if v not in (None, "")]
    if not vals:
        return None
    if all(is_formula(v) for v in vals):  # only formulas: go by header
        if any(w in h for w in PERCENT_WORDS):
            return "percent"
        return "currency" if any(w in h for w in CURRENCY_WORDS) else None
    if all(isinstance(v, datetime) for v in vals):
        return "datetime" if any(v.hour or v.minute for v in vals) else "date"
    if all(isinstance(v, date) for v in vals):
        return "date"
    nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
    if len(nums) < len(vals) * 0.9:
        return None
    if any(w in h for w in PERCENT_WORDS) and all(-1.5 <= v <= 1.5 for v in nums):
        return "percent"
    if h in PLAIN_NUMBER_WORDS or any(h.endswith(w) for w in (" id", " year", " code")):
        return "0"
    if any(w in h for w in CURRENCY_WORDS):
        return "currency" if any(isinstance(v, float) and not float(v).is_integer() for v in nums) else "currency0"
    return "integer" if all(float(v).is_integer() for v in nums) else "number"


# --------------------------------------------------------------------------- #
# Small utilities
# --------------------------------------------------------------------------- #
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
ISO_DT = re.compile(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2})?$")


def clean_str(v: str) -> str:
    return _ILLEGAL.sub("", v)[:32000]


def coerce(v: Any) -> Any:
    """ISO date strings -> real dates; strip illegal characters."""
    if isinstance(v, str):
        s = v.strip()
        try:
            if ISO_DATE.match(s):
                return date.fromisoformat(s)
            if ISO_DT.match(s):
                return datetime.fromisoformat(s.replace(" ", "T"))
        except ValueError:
            pass
        return clean_str(v)
    if isinstance(v, (dict, list)):
        return clean_str(str(v))
    return v


def cell_value(v: Any, row: int) -> Any:
    v = coerce(v)
    return v.replace("{r}", str(row)) if is_formula(v) else v


def q(sheet: str) -> str:
    """Quote a sheet name for use inside a formula."""
    return "'" + sheet.replace("'", "''") + "'"


def fill_of(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


def side_of(color: str, style: str = "thin") -> Side:
    return Side(style=style, color=color)


def safe_sheet_name(name: Any, existing: List[str]) -> str:
    n = re.sub(r"[\[\]:*?/\\]", "-", str(name or "Sheet")).strip("'").strip()[:31] or "Sheet"
    taken = {e.lower() for e in existing}
    base, i = n, 2
    while n.lower() in taken:
        suffix = f" ({i})"
        n, i = base[:31 - len(suffix)] + suffix, i + 1
    return n


def parse_cell(ref: str) -> Tuple[int, int]:
    """'B2' -> (row, col)."""
    m = re.match(r"^\$?([A-Za-z]{1,3})\$?(\d+)$", str(ref).strip())
    if not m:
        raise ExcelError(f"'{ref}' is not a valid cell reference (example: B2).", code="INVALID_REFERENCE")
    return int(m.group(2)), column_index_from_string(m.group(1).upper())


def parse_range(ws: Any, rng: str) -> Tuple[int, int, int, int]:
    """'A1:C9' -> (min_row, min_col, max_row, max_col). Handles whole columns/rows."""
    try:
        c1, r1, c2, r2 = range_boundaries(str(rng).replace("$", ""))
    except Exception:
        raise ExcelError(f"'{rng}' is not a valid range (example: A1:D20).", code="INVALID_REFERENCE")
    r1, r2 = r1 or 1, r2 or max(ws.max_row, 1)
    c1, c2 = c1 or 1, c2 or max(ws.max_column, 1)
    return r1, c1, r2, c2


def range_cells(ws: Any, rng: str):
    r1, c1, r2, c2 = parse_range(ws, rng)
    if (r2 - r1 + 1) * (c2 - c1 + 1) > MAX_CELLS_TOUCHED:
        raise ExcelError("That range is too large for one operation. Use a smaller range.", code="RANGE_TOO_LARGE")
    for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
        for c in row:
            yield c


# --------------------------------------------------------------------------- #
# Context + regions (tables registered by name so charts/summaries can refer to them)
# --------------------------------------------------------------------------- #
@dataclass
class Region:
    sheet: str
    header_row: int
    first_col: int
    last_col: int
    first_row: int
    last_row: int
    headers: List[str]
    formats: Dict[str, str] = field(default_factory=dict)
    total_row: Optional[int] = None  # a trailing Total row that sits just below the data (excluded from it)

    def col(self, ref: Any) -> int:
        """Absolute column index from a header name, column letter or index."""
        if isinstance(ref, int):
            return ref
        text = str(ref).strip()
        for i, h in enumerate(self.headers):
            if h.strip().lower() == text.lower():
                return self.first_col + i
        if re.fullmatch(r"[A-Za-z]{1,3}", text):
            idx = column_index_from_string(text.upper())
            if self.first_col <= idx <= self.last_col:
                return idx
        raise ExcelError(f"Column '{ref}' was not found. Available columns: {', '.join(self.headers)}.",
                         status=422, code="COLUMN_NOT_FOUND")

    def rng(self, ref: Any, with_header: bool = False) -> str:
        c = get_column_letter(self.col(ref))
        top = self.header_row if with_header else self.first_row
        return f"{q(self.sheet)}!${c}${top}:${c}${self.last_row}"

    @property
    def data_ref(self) -> str:
        return (f"{get_column_letter(self.first_col)}{self.header_row}:"
                f"{get_column_letter(self.last_col)}{self.last_row}")


class Ctx:
    def __init__(self, wb: Workbook, theme: Theme):
        self.wb, self.theme = wb, theme
        self.log: List[str] = []
        self.warnings: List[str] = []
        self.tables: Dict[str, Region] = {}
        self.widths: Dict[str, Dict[int, float]] = {}
        self.report: Dict[str, Any] = {}
        self.table_seq = 0
        self.deferred: List[Tuple[Any, str]] = []
        self.deferred_charts: List[Tuple[Any, Dict[str, Any], int, int]] = []

    def note(self, msg: str) -> None:
        self.log.append(msg)

    def warn(self, msg: str) -> None:
        if msg not in self.warnings:
            self.warnings.append(msg)

    def sheet(self, name: Any) -> Any:
        if name is None and len(self.wb.sheetnames) == 1:
            return self.wb[self.wb.sheetnames[0]]
        for s in self.wb.sheetnames:
            if s.lower() == str(name or "").strip().lower():
                return self.wb[s]
        raise ExcelError(f"Sheet '{name}' was not found. Sheets in this workbook: {', '.join(self.wb.sheetnames)}.",
                         status=422, code="SHEET_NOT_FOUND")

    def want_width(self, ws: Any, col: int, width: float) -> None:
        d = self.widths.setdefault(ws.title, {})
        d[col] = max(d.get(col, 0), width)

    def apply_widths(self) -> None:
        for title, d in self.widths.items():
            if title not in self.wb.sheetnames:
                continue
            ws = self.wb[title]
            for col, w in d.items():
                letter = get_column_letter(col)
                cur = ws.column_dimensions[letter].width if letter in ws.column_dimensions else None
                ws.column_dimensions[letter].width = max(6, min(max(w, cur or 0), 60))


# --------------------------------------------------------------------------- #
# Conditional formatting and data validation (shared by create + edit)
# --------------------------------------------------------------------------- #
COLOR_ALIASES = {
    "green": ("C6EFCE", "006100"), "red": ("FFC7CE", "9C0006"), "amber": ("FFEB9C", "9C5700"),
    "yellow": ("FFEB9C", "9C5700"), "blue": ("DDEBF7", "1F4E78"), "grey": ("EDEDED", "404040"),
    "gray": ("EDEDED", "404040"), "orange": ("FCE4D6", "833C0B"),
}
SCALE_PALETTES = {
    "rag": ("F8696B", "FFEB84", "63BE7B"), "gyr": ("63BE7B", "FFEB84", "F8696B"),
    "blue": ("FFFFFF", None, "5B9BD5"), "green": ("FFFFFF", None, "63BE7B"),
}
CELL_OPS = {"greaterthan": "greaterThan", ">": "greaterThan", "lessthan": "lessThan", "<": "lessThan",
            "greaterthanorequal": "greaterThanOrEqual", ">=": "greaterThanOrEqual",
            "lessthanorequal": "lessThanOrEqual", "<=": "lessThanOrEqual", "equal": "equal", "=": "equal",
            "==": "equal", "notequal": "notEqual", "!=": "notEqual", "between": "between",
            "notbetween": "notBetween"}


def colors_for(spec: Dict[str, Any], default: str = "amber") -> Tuple[str, str]:
    name = str(spec.get("color") or spec.get("fill") or default).lower()
    bg, fg = COLOR_ALIASES.get(name, (hex6(name, "FFEB9C"), "000000"))
    return hex6(spec.get("fill"), bg) if spec.get("fill") else bg, hex6(spec.get("fontColor"), fg)


def cf_literal(v: Any) -> str:
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return str(v)
    s = str(v)
    if is_formula(s):
        return s[1:]
    if ISO_DATE.match(s):
        y, m, d = s.split("-")
        return f"DATE({int(y)},{int(m)},{int(d)})"
    return '"' + s.replace('"', '""') + '"'


def add_cf(ws: Any, rng: str, spec: Dict[str, Any], theme: Theme) -> None:
    """Add one conditional-format rule. spec.type: colorScale | dataBar | iconSet | cell | text | formula |
    topN | bottomN | duplicates | blanks | status"""
    t = re.sub(r"[^a-z]", "", str(spec.get("type", "")).lower())
    tl = rng.split(":")[0].replace("$", "")
    cf = ws.conditional_formatting
    if t in ("colorscale", "scale"):
        lo, mid, hi = SCALE_PALETTES.get(str(spec.get("palette", "rag")).lower(), SCALE_PALETTES["rag"])
        if mid:
            cf.add(rng, ColorScaleRule(start_type="min", start_color=lo, mid_type="percentile", mid_value=50,
                                       mid_color=mid, end_type="max", end_color=hi))
        else:
            cf.add(rng, ColorScaleRule(start_type="min", start_color=lo, end_type="max", end_color=hi))
    elif t == "databar":
        cf.add(rng, DataBarRule(start_type="min", end_type="max", color=hex6(spec.get("color"), theme.primary),
                                showValue=True))
    elif t == "iconset":
        name = spec.get("icons", "3TrafficLights1")
        cf.add(rng, IconSetRule(name, "percent", [0, 33, 67], showValue=None))
    elif t in ("cell", "cellis", "value"):
        op = CELL_OPS.get(str(spec.get("operator", "greaterThan")).lower().replace(" ", ""), "greaterThan")
        vals = spec.get("values") if op in ("between", "notBetween") else [spec.get("value")]
        if not vals or any(v is None for v in vals):
            raise ExcelError("A 'cell' conditional format needs 'value' (or 'values' for between).", code="INVALID_CF")
        bg, fg = colors_for(spec)
        cf.add(rng, CellIsRule(operator=op, formula=[cf_literal(v) for v in vals], fill=fill_of(bg),
                               font=Font(color=fg, bold=bool(spec.get("bold")))))
    elif t in ("text", "contains"):
        bg, fg = colors_for(spec)
        txt = str(spec.get("text", "")).replace('"', '""')
        cf.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{txt}",{tl}))'], fill=fill_of(bg), font=Font(color=fg)))
    elif t == "formula":
        bg, fg = colors_for(spec)
        f = str(spec.get("formula", "")).lstrip("=")
        if not f:
            raise ExcelError("A 'formula' conditional format needs 'formula'.", code="INVALID_CF")
        cf.add(rng, FormulaRule(formula=[f], fill=fill_of(bg), font=Font(color=fg)))
    elif t in ("topn", "bottomn", "top", "bottom"):
        bg, fg = colors_for(spec, "green" if t.startswith("top") else "red")
        dxf = DifferentialStyle(fill=fill_of(bg), font=Font(color=fg))
        cf.add(rng, Rule(type="top10", rank=int(spec.get("n", 10)), bottom=t.startswith("bottom"),
                         percent=bool(spec.get("percent")), dxf=dxf))
    elif t in ("duplicates", "duplicate"):
        bg, fg = colors_for(spec, "red")
        cf.add(rng, Rule(type="duplicateValues", dxf=DifferentialStyle(fill=fill_of(bg), font=Font(color=fg))))
    elif t == "blanks":
        bg, fg = colors_for(spec, "amber")
        cf.add(rng, FormulaRule(formula=[f"LEN(TRIM({tl}))=0"], fill=fill_of(bg), font=Font(color=fg)))
    elif t == "status":  # {"Done": "green", "Late": "red"}
        for word, color in (spec.get("map") or {}).items():
            bg, fg = COLOR_ALIASES.get(str(color).lower(), (hex6(color, "FFEB9C"), "000000"))
            cf.add(rng, CellIsRule(operator="equal", formula=[cf_literal(str(word))], fill=fill_of(bg),
                                   font=Font(color=fg, bold=True)))
    else:
        raise ExcelError(f"Unknown conditional format type '{spec.get('type')}'. Use colorScale, dataBar, iconSet, "
                         "cell, text, formula, topN, bottomN, duplicates, blanks or status.", code="INVALID_CF")


def add_validation(ctx: Ctx, ws: Any, rng: str, spec: Dict[str, Any]) -> None:
    """spec.type: list | whole | decimal | date | textLength | custom"""
    t = re.sub(r"[^a-z]", "", str(spec.get("type", "list")).lower())
    kind = {"list": "list", "whole": "whole", "integer": "whole", "decimal": "decimal", "number": "decimal",
            "date": "date", "textlength": "textLength", "length": "textLength", "custom": "custom"}.get(t)
    if not kind:
        raise ExcelError(f"Unknown validation type '{spec.get('type')}'.", code="INVALID_VALIDATION")
    f1: Any = None
    f2: Any = None
    op = None
    if kind == "list":
        if spec.get("source"):
            f1 = str(spec["source"]) if str(spec["source"]).startswith("=") else "=" + str(spec["source"])
            f1 = f1[1:] if f1.startswith("=") else f1
        else:
            values = [str(v).replace(",", ";") for v in spec.get("values", [])]
            if not values:
                raise ExcelError("A list validation needs 'values' or 'source'.", code="INVALID_VALIDATION")
            joined = ",".join(values)
            if len(joined) > 250:  # inline lists are limited to 255 characters: park the list on a hidden sheet
                lst = ctx.wb["_Lists"] if "_Lists" in ctx.wb.sheetnames else ctx.wb.create_sheet("_Lists")
                lst.sheet_state = "hidden"
                col = lst.max_column + (1 if lst.max_row > 1 or lst["A1"].value else 0)
                for i, v in enumerate(values, 1):
                    lst.cell(i, col, v)
                L = get_column_letter(col)
                f1 = f"'_Lists'!${L}$1:${L}${len(values)}"
            else:
                f1 = f'"{joined}"'
    elif kind == "custom":
        f1 = str(spec.get("formula", "")).lstrip("=")
    else:
        op = {"between": "between", "notbetween": "notBetween", "equal": "equal", "greaterthan": "greaterThan",
              "lessthan": "lessThan", "greaterthanorequal": "greaterThanOrEqual",
              "lessthanorequal": "lessThanOrEqual"}.get(str(spec.get("operator", "between")).lower(), "between")
        lo, hi = spec.get("min", spec.get("value")), spec.get("max")
        if op == "between" and (lo is None or hi is None):
            raise ExcelError("A 'between' validation needs 'min' and 'max'.", code="INVALID_VALIDATION")
        f1 = cf_literal(lo) if kind == "date" else str(lo)
        f2 = (cf_literal(hi) if kind == "date" else str(hi)) if hi is not None else None
    dv = DataValidation(type=kind, operator=op, formula1=f1, formula2=f2, allow_blank=spec.get("allowBlank", True),
                        showErrorMessage=spec.get("showError", True), showInputMessage=bool(spec.get("prompt")),
                        errorStyle={"warning": "warning", "info": "information"}.get(
                            str(spec.get("errorStyle", "stop")).lower(), "stop"))
    dv.errorTitle = str(spec.get("errorTitle", "Invalid entry"))[:32]
    dv.error = str(spec.get("error", "The value entered is not allowed."))[:255]
    if spec.get("prompt"):
        dv.promptTitle = str(spec.get("promptTitle", "Input"))[:32]
        dv.prompt = str(spec["prompt"])[:255]
    dv.add(rng)
    ws.add_data_validation(dv)


# --------------------------------------------------------------------------- #
# Block renderers
# --------------------------------------------------------------------------- #
def merge_fill(ws: Any, r1: int, c1: int, r2: int, c2: int, color: Optional[str]) -> None:
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    if color:
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(r, c).fill = fill_of(color)


def outline(ws: Any, r1: int, c1: int, r2: int, c2: int, color: str) -> None:
    s = side_of(color)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            cell.border = Border(left=s if c == c1 else None, right=s if c == c2 else None,
                                 top=s if r == r1 else None, bottom=s if r == r2 else None)


def block_title(ctx: Ctx, ws: Any, b: Dict[str, Any], c0: int, r: int, span: int) -> int:
    th = ctx.theme
    merge_fill(ws, r, c0, r, c0 + span - 1, th.dark)
    c = ws.cell(r, c0, coerce(str(b.get("text") or b.get("title") or ws.title)))
    c.font = Font(name=th.font, size=20, bold=True, color="FFFFFF")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[r].height = 38
    r += 1
    if b.get("subtitle"):
        merge_fill(ws, r, c0, r, c0 + span - 1, th.primary)
        c = ws.cell(r, c0, coerce(str(b["subtitle"])))
        c.font = Font(name=th.font, size=11, italic=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 22
        r += 1
    return r + 1


def block_text(ctx: Ctx, ws: Any, b: Dict[str, Any], c0: int, r: int, span: int) -> int:
    th = ctx.theme
    style = str(b.get("style", "normal")).lower()
    lines = b.get("lines") or ([b["text"]] if b.get("text") else [])
    width_chars = max(span * 12, 40)
    for line in lines:
        text = coerce(str(line))
        if style == "bullets":
            text = "•  " + str(text)
        ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + span - 1)
        c = ws.cell(r, c0, text)
        size, bold, italic, color = {"heading": (14, True, False, th.dark), "note": (10, False, True, th.muted),
                                     "bullets": (11, False, False, th.text)}.get(style, (11, False, False, th.text))
        c.font = Font(name=th.font, size=size, bold=bold, italic=italic, color=color)
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        ws.row_dimensions[r].height = max(18, 15 * math.ceil(len(str(text)) / width_chars) + 3)
        r += 1
    return r + 1


def block_kpis(ctx: Ctx, ws: Any, b: Dict[str, Any], c0: int, r: int, span: int) -> int:
    th = ctx.theme
    items = b.get("items") or b.get("kpis") or []
    per_row = max(1, (span + 1) // 3)
    for i, k in enumerate(items):
        row_block, pos = divmod(i, per_row)
        r0, col = r + row_block * 5, c0 + pos * 3
        for cc in (col, col + 1):
            ctx.want_width(ws, cc, 14)
        merge_fill(ws, r0, col, r0, col + 1, th.light)
        lab = ws.cell(r0, col, coerce(str(k.get("label", "")).upper()))
        lab.font = Font(name=th.font, size=9, bold=True, color=th.muted)
        lab.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        merge_fill(ws, r0 + 1, col, r0 + 2, col + 1, th.light)
        val = ws.cell(r0 + 1, col, cell_value(k.get("value"), r0 + 1))
        if is_formula(val.value):
            ctx.deferred.append((ws, val.coordinate))
        val.font = Font(name=th.font, size=24, bold=True, color=hex6(k.get("color"), th.accents[i % len(th.accents)]))
        val.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        fmt = number_format(k.get("format"), th)
        if fmt:
            val.number_format = fmt
        merge_fill(ws, r0 + 3, col, r0 + 3, col + 1, th.light)
        note = ws.cell(r0 + 3, col, coerce(str(k.get("note", ""))))
        note.font = Font(name=th.font, size=9, italic=True, color=th.muted)
        note.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        outline(ws, r0, col, r0 + 3, col + 1, th.border)
        ws.row_dimensions[r0].height = 20
        ws.row_dimensions[r0 + 1].height = 20
        ws.row_dimensions[r0 + 2].height = 20
    rows = math.ceil(len(items) / per_row) * 5 if items else 0
    return r + rows + 1


import json as _json  # noqa: E402

HEADER_KEYS = ("headers", "header", "columnNames", "column_names", "cols", "fields", "columns")
ROW_KEYS = ("rows", "data", "values", "records", "items", "body")
NAME_KEYS = ("header", "name", "title", "label", "field", "key", "text")
TYPE_TO_FORMAT = {"date": "date", "datetime": "datetime", "currency": "currency", "money": "currency",
                  "percent": "percent", "percentage": "percent", "number": "number", "decimal": "number",
                  "integer": "integer", "int": "integer", "text": None, "string": None}


def _maybe_json(v: Any) -> Any:
    """Copilot Studio sometimes sends arrays/objects as JSON text."""
    if isinstance(v, str) and v.strip()[:1] in ("[", "{"):
        try:
            return _json.loads(v)
        except ValueError:
            return v
    return v


def _col_name(c: Any) -> Optional[str]:
    if isinstance(c, dict):
        for k in NAME_KEYS:
            if c.get(k) not in (None, ""):
                return str(c[k])
        return None
    return str(c) if c not in (None, "") else None


def normalize_table(b: Dict[str, Any]) -> Tuple[List[str], List[List[Any]], List[Dict[str, Any]]]:
    """Accept the many shapes an agent may send for a table; return (headers, rows, per-column specs)."""
    b = dict(b)
    for k in ROW_KEYS:  # nested form: {"data": {"headers": [...], "rows": [...]}}
        inner = _maybe_json(b.get(k))
        if isinstance(inner, dict):
            for kk in HEADER_KEYS + ROW_KEYS:
                if kk in inner and not b.get(kk):
                    b[kk] = inner[kk]
            b.pop(k, None) if k not in inner else None
    raw_cols = _maybe_json(b.get("columns"))
    raw_cols = raw_cols if isinstance(raw_cols, list) else []
    rows: Any = next((v for v in (_maybe_json(b.get(k)) for k in ROW_KEYS) if isinstance(v, list)), [])
    raw_headers: Any = None
    for k in HEADER_KEYS:
        v = _maybe_json(b.get(k))
        if isinstance(v, str) and v.strip():
            v = [h.strip() for h in v.split(",")]
        if isinstance(v, list) and v:
            raw_headers = v
            break
    headers: Optional[List[Optional[str]]] = [_col_name(c) for c in raw_headers] if raw_headers else None
    rows = [_maybe_json(r) for r in rows]
    if rows and all(isinstance(r, dict) for r in rows):
        if not headers:
            headers = list(OrderedDict.fromkeys(k for row in rows for k in row))
        lowered = [{str(k).lower(): v for k, v in row.items()} for row in rows]
        rows = [[row.get(h) if h in row else lw.get(str(h).lower()) for h in headers] for row, lw in zip(rows, lowered)]
    elif rows and not all(isinstance(r, (list, tuple)) for r in rows):
        rows = [[r] for r in rows]  # a flat list = one column
    if not headers and rows:  # no header given: first row is the header when it is all text
        first = rows[0]
        if all(isinstance(v, str) and v.strip() for v in first) and len(rows) >= 1:
            headers, rows = list(first), rows[1:]
    if headers and rows and isinstance(rows[0], (list, tuple)) and [str(v).strip().lower() for v in rows[0]] == \
            [str(h).strip().lower() for h in headers]:
        rows = rows[1:]  # the agent repeated the header as the first data row
    if not headers:
        raise ExcelError("A table needs 'headers' (a list of column names) and 'rows' (a list of row lists). "
                         f"This table block had these fields: {', '.join(sorted(map(str, b.keys()))) or 'none'}. Example: "
                         "{\"type\":\"table\",\"name\":\"Sales\",\"headers\":[\"Region\",\"Revenue\"],"
                         "\"rows\":[[\"North\",1200],[\"South\",900]]}", status=422, code="INVALID_TABLE")
    headers = [clean_str(str(h)) if h not in (None, "") else f"Column {i + 1}" for i, h in enumerate(headers)]
    if len(rows) > MAX_ROWS_IN_SPEC:
        raise ExcelError(f"Too many rows ({len(rows)}); the limit is {MAX_ROWS_IN_SPEC}.", code="TOO_MANY_ROWS")
    n = len(headers)
    rows = [list(r)[:n] + [None] * (n - len(r)) for r in rows]
    specs: List[Dict[str, Any]] = []
    for i, h in enumerate(headers):
        spec: Dict[str, Any] = {}
        cand = raw_cols[i] if i < len(raw_cols) and isinstance(raw_cols[i], dict) and _col_name(raw_cols[i]) in (h, None) else None
        if cand is None:
            cand = next((c for c in raw_cols if isinstance(c, dict) and (_col_name(c) or "").lower() == h.lower()), None)
        if cand:
            spec = dict(cand)
            kind = str(spec.get("type") or spec.get("dataType") or "").lower()
            if not spec.get("format") and kind in TYPE_TO_FORMAT and TYPE_TO_FORMAT[kind]:
                spec["format"] = TYPE_TO_FORMAT[kind]
        specs.append(spec)
    return headers, rows, specs


def block_table(ctx: Ctx, ws: Any, b: Dict[str, Any], c0: int, r: int, span: int) -> int:
    th = ctx.theme
    headers, rows, specs = normalize_table(b)
    n = len(headers)

    ctx.table_seq += 1
    name = str(b.get("name") or f"Table{ctx.table_seq}")
    if b.get("title"):
        c = ws.cell(r, c0, coerce(str(b["title"])))
        c.font = Font(name=th.font, size=13, bold=True, color=th.dark)
        r += 1
    header_row, first_row = r, r + 1
    letters = {h.lower(): get_column_letter(c0 + i) for i, h in enumerate(headers)}
    fmts: Dict[str, str] = {}
    ctx.tables[name.lower()] = Region(ws.title, header_row, c0, c0 + n - 1, first_row,
                                      first_row + len(rows) - 1, headers, fmts)  # registered early: self-references work

    # formula columns: fill blanks with the column template
    for ci, spec in enumerate(specs):
        if spec.get("formula"):
            for row in rows:
                if row[ci] in (None, ""):
                    row[ci] = spec["formula"]

    for ci, h in enumerate(headers):
        col = c0 + ci
        spec = specs[ci]
        key = spec.get("format") or infer_format(h, [coerce(row[ci]) for row in rows[:500]])
        fmt = number_format(key, th) if key else None
        if fmt and fmt not in ("General",):
            fmts[h] = fmt
        hc = ws.cell(header_row, col, h)
        hc.font = Font(name=th.font, size=11, bold=True, color="FFFFFF")
        hc.fill = fill_of(th.dark)
        hc.alignment = Alignment(horizontal=spec.get("align", "center"), vertical="center", wrap_text=True)
        hc.border = Border(bottom=side_of(th.primary, "medium"))
        sample = [len(f"{v:,.2f}") + 2 if isinstance(v, (int, float)) and not isinstance(v, bool)
                  else (11 if isinstance(v, (date, datetime)) else len(str(v))) for v in
                  [row[ci] for row in rows[:300] if row[ci] is not None and not is_formula(row[ci])]]
        if any(is_formula(row[ci]) for row in rows[:50]) or spec.get("total"):
            sample.append(14)  # formula results/totals are wider than their text suggests
        width = spec.get("width") or max(len(h) + 3, min(max(sample or [8]) + 2, 48))
        ctx.want_width(ws, col, float(width))
    ws.row_dimensions[header_row].height = 28

    numeric_like = {h for h in headers if h in fmts and fmts[h] != "@"}
    for ri, row in enumerate(rows):
        rr = first_row + ri
        banded = bool(b.get("banding", True)) and ri % 2 == 1
        for ci, h in enumerate(headers):
            spec = specs[ci]
            cell = ws.cell(rr, c0 + ci, resolve_tokens(ctx, cell_value(row[ci], rr), letters, rr))
            cell.font = Font(name=th.font, size=11, color=th.text)
            if h in fmts:
                cell.number_format = fmts[h]
            is_num = isinstance(cell.value, (int, float, date, datetime)) or is_formula(cell.value) and h in numeric_like
            horiz = spec.get("align") or ("right" if is_num and not isinstance(cell.value, (date, datetime)) else
                                          "center" if isinstance(cell.value, (date, datetime)) else "left")
            cell.alignment = Alignment(horizontal=horiz, vertical="center", wrap_text=bool(spec.get("wrap")))
            if banded:
                cell.fill = fill_of(th.band)
            cell.border = Border(bottom=side_of(th.border))
    last_row = first_row + len(rows) - 1

    # totals row (SUBTOTAL ignores filtered-out rows)
    total_row = None
    totals = {ci: s.get("total") for ci, s in enumerate(specs) if s.get("total")}
    if b.get("totalRow") or totals:
        total_row = last_row + 1
        codes = {"sum": 109, "avg": 101, "average": 101, "count": 102, "counta": 103, "max": 104, "min": 105}
        for ci, h in enumerate(headers):
            col = c0 + ci
            cell = ws.cell(total_row, col)
            agg = totals.get(ci)
            L = get_column_letter(col)
            if agg and last_row >= first_row:
                a = str(agg).lower()
                cell.value = (f"=SUBTOTAL({codes[a]},{L}{first_row}:{L}{last_row})" if a in codes
                              else (resolve_tokens(ctx, str(agg), letters, last_row) if is_formula(agg) else None))
            elif ci == 0 and not agg:
                cell.value = "Total"
            cell.font = Font(name=th.font, size=11, bold=True, color=th.dark)
            cell.fill = fill_of(th.light)
            cell.border = Border(top=side_of(th.dark, "medium"))
            if h in fmts:
                cell.number_format = fmts[h]
            cell.alignment = Alignment(horizontal="right" if ci else "left", vertical="center")

    # column-level conditional formats + validation
    if last_row >= first_row:
        for ci, h in enumerate(headers):
            L = get_column_letter(c0 + ci)
            rng = f"{L}{first_row}:{L}{last_row}"
            cfs = specs[ci].get("cf") or specs[ci].get("conditionalFormat")
            for cf in (cfs if isinstance(cfs, list) else [cfs] if cfs else []):
                add_cf(ws, rng, cf, th)
            if specs[ci].get("validation"):
                add_validation(ctx, ws, f"{L}{first_row}:{L}{max(last_row, first_row) + int(b.get('validationExtraRows', 200))}",
                               specs[ci]["validation"])
    for cf in b.get("conditionalFormats", []) or []:  # table-level rules referencing a column by header
        L = get_column_letter(c0 + headers.index(cf["column"])) if cf.get("column") in headers else None
        if L:
            add_cf(ws, f"{L}{first_row}:{L}{last_row}", cf, th)

    # structure: filter or Excel Table object
    ref = f"{get_column_letter(c0)}{header_row}:{get_column_letter(c0 + n - 1)}{max(last_row, first_row)}"
    if b.get("excelTable"):
        tab = Table(displayName=re.sub(r"[^A-Za-z0-9_]", "_", name)[:60] or f"Table{ctx.table_seq}", ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name=b.get("tableStyle", "TableStyleMedium2"), showRowStripes=True)
        ws.add_table(tab)
    elif b.get("filter", True) and not ws.auto_filter.ref and total_row is None:
        ws.auto_filter.ref = ref
    if b.get("freezeHeader") or (b.get("freezeHeader") is None and not ws.freeze_panes and header_row <= 12
                                 and len(rows) > 12):
        ws.freeze_panes = ws.cell(first_row, 1)

    ctx.note(f"Table '{name}' on '{ws.title}' ({len(rows)} rows, {n} columns)")
    return (total_row or last_row) + 2


TOKEN_ROW = re.compile(r"\[@([^\]]+)\]")
TOKEN_COL = re.compile(r"\b([A-Za-z_][A-Za-z0-9_]*)\[([^\]@][^\]]*)\]")


def resolve_tokens(ctx: "Ctx", f: Any, letters: Optional[Dict[str, str]] = None, row: Optional[int] = None) -> Any:
    """Friendly formula syntax so the agent never has to know column letters:
       [@Units]*[@Price]    -> same-row cells of the current table (e.g. E5*F5)
       SUM(Sales[Revenue])  -> absolute range of that column of table 'Sales'"""
    if not is_formula(f):
        return f
    if letters and row is not None:
        f = TOKEN_ROW.sub(lambda m: (f"{letters[m.group(1).strip().lower()]}{row}"
                                     if m.group(1).strip().lower() in letters else m.group(0)), f)

    def col_sub(m: "re.Match[str]") -> str:
        reg = ctx.tables.get(m.group(1).lower())
        if not reg:
            return m.group(0)
        try:
            return reg.rng(m.group(2).strip())
        except ExcelError:
            return m.group(0)

    return TOKEN_COL.sub(col_sub, f)


def finalize_tokens(ctx: "Ctx") -> None:
    """Resolve Table[Column] tokens in formulas written before their table existed (e.g. KPI cards)."""
    for ws, coord in ctx.deferred:
        cell = ws[coord]
        new = resolve_tokens(ctx, cell.value)
        if new != cell.value:
            cell.value = new
        elif is_formula(cell.value) and TOKEN_COL.search(cell.value or "") and "!" not in cell.value:
            ctx.warn(f"Formula in {ws.title}!{coord} references a table that does not exist: {cell.value}")
    ctx.deferred.clear()
    pending, ctx.deferred_charts = ctx.deferred_charts, []
    for ws, b, c0, r in pending:
        reg = lookup_region(ctx, (b.get("data") or {}).get("table") or b.get("table") or b.get("source"))
        chart, _ = build_chart(ctx, b, reg)
        place = str(b.get("place", "below"))
        ws.add_chart(chart, place.upper() if re.fullmatch(r"[A-Za-z]{1,3}\d+", place) else f"{get_column_letter(c0)}{r}")


def lookup_region(ctx: Ctx, ref: Any) -> Region:
    key = str(ref or "").lower()
    if key in ctx.tables:
        return ctx.tables[key]
    if not ref and len(ctx.tables) == 1:
        return next(iter(ctx.tables.values()))
    raise ExcelError(f"Table '{ref}' was not found. Known tables: {', '.join(ctx.tables) or 'none'}.",
                     status=422, code="TABLE_NOT_FOUND")


def pyvalues(ws: Any, reg: Region, header: Any) -> List[Any]:
    c = reg.col(header)
    return [ws.cell(r, c).value for r in range(reg.first_row, reg.last_row + 1)]


AGG_FORMULAS = {"sum": "SUMIFS({v},{g},{k})", "count": "COUNTIFS({g},{k})", "avg": "IFERROR(AVERAGEIFS({v},{g},{k}),0)",
                "average": "IFERROR(AVERAGEIFS({v},{g},{k}),0)", "min": "MINIFS({v},{g},{k})", "max": "MAXIFS({v},{g},{k})"}


def block_summary(ctx: Ctx, ws: Any, b: Dict[str, Any], c0: int, r: int, span: int) -> int:
    """Pivot-style summary built from SUMIFS/COUNTIFS formulas (live, but groups are fixed at build time)."""
    th = ctx.theme
    src = lookup_region(ctx, b.get("source"))
    src_ws = ctx.sheet(src.sheet)
    group = b.get("groupBy")
    if not group:
        raise ExcelError("A summary block needs 'groupBy' (a column header of the source table).", code="INVALID_SUMMARY")
    gcol = src.col(group)
    gname = src.headers[gcol - src.first_col]
    raw = pyvalues(src_ws, src, group)
    groups = list(OrderedDict.fromkeys(str(v).strip() if v is not None else "(blank)" for v in raw
                                       if v is not None and not is_formula(v)))
    values = b.get("values") or [{"column": None, "agg": "count"}]
    # sort by first numeric aggregate (computed in Python when the source holds plain numbers)
    first = values[0]
    sortdir = str(b.get("sort", "desc")).lower()
    if first.get("column") and first.get("agg", "sum") in ("sum", "avg", "average"):
        nums = pyvalues(src_ws, src, first["column"])
        agg: Dict[str, List[float]] = {}
        ok = True
        for gv, nv in zip(raw, nums):
            if isinstance(nv, (int, float)) and not isinstance(nv, bool):
                agg.setdefault(str(gv).strip(), []).append(float(nv))
            elif is_formula(nv):
                ok = False
        if ok and agg:
            fn = (lambda xs: sum(xs)) if first.get("agg", "sum") == "sum" else (lambda xs: sum(xs) / len(xs))
            groups.sort(key=lambda g: fn(agg.get(g, [0])), reverse=(sortdir != "asc"))
        elif sortdir in ("asc", "desc"):
            groups.sort(key=str.lower, reverse=(sortdir == "desc") and False)
    elif sortdir in ("asc", "desc"):
        groups.sort(key=str.lower)
    if b.get("top"):
        groups = groups[: int(b["top"])]
    name = str(b.get("name") or f"Summary{ctx.table_seq + 1}")
    ctx.table_seq += 1

    if b.get("title"):
        c = ws.cell(r, c0, coerce(str(b["title"])))
        c.font = Font(name=th.font, size=13, bold=True, color=th.dark)
        r += 1
    header_row, first_row = r, r + 1
    out_headers = [gname]
    for v in values:
        agg_name = str(v.get("agg", "sum")).lower()
        out_headers.append(v.get("label") or (f"{agg_name.title()} of {v['column']}" if v.get("column") else "Count"))
    if b.get("percentOfTotal"):
        out_headers.append("% of total")
    fmts: Dict[str, str] = {}
    for ci, h in enumerate(out_headers):
        c = ws.cell(header_row, c0 + ci, h)
        c.font = Font(name=th.font, size=11, bold=True, color="FFFFFF")
        c.fill = fill_of(th.dark)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ctx.want_width(ws, c0 + ci, max(16 if ci else 20, len(h) + 3))
    ws.row_dimensions[header_row].height = 28
    g_rng = src.rng(group)
    for ri, g in enumerate(groups):
        rr = first_row + ri
        key_cell = ws.cell(rr, c0, g)
        key_cell.font = Font(name=th.font, size=11, color=th.text)
        key_cell.alignment = Alignment(horizontal="left")
        for vi, v in enumerate(values):
            agg_name = str(v.get("agg", "sum")).lower()
            if agg_name not in AGG_FORMULAS:
                raise ExcelError(f"Unknown aggregation '{agg_name}'. Use sum, count, avg, min or max.", code="INVALID_SUMMARY")
            if agg_name != "count" and not v.get("column"):
                raise ExcelError(f"Aggregation '{agg_name}' needs a 'column'.", code="INVALID_SUMMARY")
            formula = "=" + AGG_FORMULAS[agg_name].format(
                v=src.rng(v["column"]) if v.get("column") else "", g=g_rng, k=f"${get_column_letter(c0)}{rr}")
            cell = ws.cell(rr, c0 + 1 + vi, formula)
            cell.font = Font(name=th.font, size=11, color=th.text)
            fmt = number_format(v.get("format"), th) or (src.formats.get(src.headers[src.col(v["column"]) - src.first_col])
                                                         if v.get("column") and agg_name != "count" else "#,##0")
            if fmt:
                cell.number_format = fmt
                fmts[out_headers[1 + vi]] = fmt
            cell.alignment = Alignment(horizontal="right")
        if b.get("percentOfTotal"):
            L = get_column_letter(c0 + 1)
            cell = ws.cell(rr, c0 + 1 + len(values), f"={L}{rr}/SUM({L}${first_row}:{L}${first_row + len(groups) - 1})")
            cell.number_format = "0.0%"
            cell.alignment = Alignment(horizontal="right")
            fmts["% of total"] = "0.0%"
        for ci in range(len(out_headers)):
            cc = ws.cell(rr, c0 + ci)
            cc.border = Border(bottom=side_of(th.border))
            if ri % 2 == 1:
                cc.fill = fill_of(th.band)
    last_row = first_row + len(groups) - 1
    total_row = None
    if b.get("totalRow", True) and groups:
        total_row = last_row + 1
        for ci in range(len(out_headers)):
            cell = ws.cell(total_row, c0 + ci)
            L = get_column_letter(c0 + ci)
            if ci == 0:
                cell.value = "Total"
            elif ci <= len(values):
                agg_name = str(values[ci - 1].get("agg", "sum")).lower()
                fn = {"sum": "SUM", "count": "SUM", "avg": "AVERAGE", "average": "AVERAGE", "min": "MIN", "max": "MAX"}[agg_name]
                cell.value = f"={fn}({L}{first_row}:{L}{last_row})"
                cell.number_format = fmts.get(out_headers[ci], "General")
            else:
                cell.value = f"=SUM({L}{first_row}:{L}{last_row})"
                cell.number_format = "0.0%"
            cell.font = Font(name=th.font, size=11, bold=True, color=th.dark)
            cell.fill = fill_of(th.light)
            cell.border = Border(top=side_of(th.dark, "medium"))
            cell.alignment = Alignment(horizontal="right" if ci else "left")
    ctx.tables[name.lower()] = Region(ws.title, header_row, c0, c0 + len(out_headers) - 1, first_row, last_row,
                                      out_headers, fmts)
    ctx.note(f"Summary '{name}' of {gname} ({len(groups)} groups)")
    return (total_row or last_row) + 2


# ---- charts ---------------------------------------------------------------- #
def numeric_headers(ws: Any, reg: Region, exclude: Optional[str] = None) -> List[str]:
    out = []
    for h in reg.headers:
        if exclude and h.lower() == exclude.lower():
            continue
        vals = [ws.cell(r, reg.col(h)).value for r in range(reg.first_row, min(reg.last_row, reg.first_row + 30) + 1)]
        vals = [v for v in vals if v not in (None, "")]
        if vals and all((isinstance(v, (int, float)) and not isinstance(v, bool)) or is_formula(v) for v in vals):
            out.append(h)
    return out


def build_chart(ctx: Ctx, spec: Dict[str, Any], reg: Region) -> Tuple[Any, int]:
    th = ctx.theme
    src_ws = ctx.sheet(reg.sheet)
    if reg.last_row < reg.first_row:
        raise ExcelError("The chart's source table has no data rows.", status=422, code="CHART_NO_DATA")
    ctype = re.sub(r"[^a-z]", "", str(spec.get("chartType", spec.get("type", "column"))).lower())
    data = spec.get("data") or {}
    cat = data.get("categories") or spec.get("categories") or reg.headers[0]
    cat_col = reg.col(cat)
    ser = data.get("series") or spec.get("series")
    if not ser:
        ser = numeric_headers(src_ws, reg, exclude=reg.headers[cat_col - reg.first_col])[:4]
    if not ser:
        raise ExcelError("No numeric columns found to chart. Name them in 'series'.", status=422, code="CHART_NO_SERIES")
    ser_cols = [reg.col(s["column"] if isinstance(s, dict) else s) for s in ser]
    cats = Reference(src_ws, min_col=cat_col, min_row=reg.first_row, max_row=reg.last_row)

    def refs(cols: List[int]):
        return [Reference(src_ws, min_col=c, min_row=reg.header_row, max_row=reg.last_row) for c in cols]

    if ctype in ("pie", "doughnut", "donut"):
        chart: Any = PieChart() if ctype == "pie" else DoughnutChart()
        chart.add_data(refs(ser_cols[:1])[0], titles_from_data=True)
        chart.set_categories(cats)
        s = chart.series[0]
        for i in range(reg.last_row - reg.first_row + 1):
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = th.accents[i % len(th.accents)]
            s.dPt.append(pt)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        for a in ("showVal", "showCatName", "showSerName", "showLegendKey"):
            setattr(chart.dataLabels, a, False)
        if ctype != "pie":
            chart.holeSize = 55
    elif ctype == "scatter":
        chart = ScatterChart()
        chart.style = 13
        xref = Reference(src_ws, min_col=cat_col, min_row=reg.first_row, max_row=reg.last_row)
        for i, ref in enumerate(refs(ser_cols)):
            s = Series(ref, xref, title_from_data=True)
            s.marker.symbol = "circle"
            s.graphicalProperties.line.noFill = True
            s.marker.graphicalProperties = GraphicalProperties(solidFill=th.accents[i % len(th.accents)])
            chart.series.append(s)
    elif ctype in ("line", "linemarkers"):
        chart = LineChart()
        for i, ref in enumerate(refs(ser_cols)):
            chart.add_data(ref, titles_from_data=True)
        chart.set_categories(cats)
        for i, s in enumerate(chart.series):
            color = th.accents[i % len(th.accents)]
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 28575
            s.marker.symbol = "circle"
            s.marker.size = 6
            s.marker.graphicalProperties = GraphicalProperties(solidFill=color)
            s.smooth = False
    elif ctype == "area":
        chart = AreaChart()
        for ref in refs(ser_cols):
            chart.add_data(ref, titles_from_data=True)
        chart.set_categories(cats)
        for i, s in enumerate(chart.series):
            s.graphicalProperties.solidFill = th.accents[i % len(th.accents)]
    else:  # column / bar / stacked / combo
        chart = BarChart()
        chart.type = "bar" if ctype in ("bar", "horizontalbar", "barh", "stackedbar") else "col"
        stacked = ctype.startswith("stacked")
        pct = ctype.startswith("percent")
        chart.grouping = "percentStacked" if pct else "stacked" if stacked else "clustered"
        if stacked or pct:
            chart.overlap = 100
        chart.gapWidth = 60
        line_cols = []
        if ctype == "combo":
            names = [str(x).lower() for x in (spec.get("lineSeries") or [])]
            line_cols = [c for c in ser_cols if reg.headers[c - reg.first_col].lower() in names]
        bar_cols = [c for c in ser_cols if c not in line_cols]
        for ref in refs(bar_cols):
            chart.add_data(ref, titles_from_data=True)
        chart.set_categories(cats)
        for i, s in enumerate(chart.series):
            s.graphicalProperties.solidFill = th.accents[i % len(th.accents)]
            s.graphicalProperties.line.noFill = True
        if line_cols:
            line = LineChart()
            for ref in refs(line_cols):
                line.add_data(ref, titles_from_data=True)
            for j, s in enumerate(line.series):
                color = th.accents[(len(bar_cols) + j) % len(th.accents)]
                s.graphicalProperties.line.solidFill = color
                s.graphicalProperties.line.width = 28575
                s.marker.symbol = "circle"
                s.marker.graphicalProperties = GraphicalProperties(solidFill=color)
                s.smooth = False
            if spec.get("secondaryAxis", True):
                line.y_axis.axId = 200
                line.y_axis.crosses = "max"
                line.y_axis.majorGridlines = None
            chart += line

    if spec.get("title"):
        chart.title = str(spec["title"])
    chart.width = float(spec.get("width", 17))
    chart.height = float(spec.get("height", 8.5))
    legend = spec.get("legend")
    if legend is False or (legend is None and len(ser_cols) == 1 and ctype not in ("pie", "doughnut", "donut")):
        chart.legend = None
    elif chart.legend is not None:
        chart.legend.position = {"right": "r", "top": "t", "left": "l"}.get(str(legend).lower(), "b")
    if ctype not in ("pie", "doughnut", "donut"):
        try:
            chart.x_axis.delete = False
            chart.y_axis.delete = False
            if ctype == "scatter":
                spec = {**spec, "xTitle": spec.get("xTitle") or reg.headers[cat_col - reg.first_col],
                        "yTitle": spec.get("yTitle") or reg.headers[ser_cols[0] - reg.first_col]}
                xf = reg.formats.get(reg.headers[cat_col - reg.first_col])
                if xf and xf != "@":
                    chart.x_axis.number_format = xf
            if spec.get("xTitle"):
                chart.x_axis.title = str(spec["xTitle"])
            if spec.get("yTitle"):
                chart.y_axis.title = str(spec["yTitle"])
            if chart.y_axis.majorGridlines is not None:
                chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=th.border))
            fmt = number_format(spec.get("axisFormat"), th) or next(
                (reg.formats[reg.headers[c - reg.first_col]] for c in ser_cols[:1]
                 if reg.headers[c - reg.first_col] in reg.formats), None)
            if fmt and fmt != "@":
                chart.y_axis.number_format = fmt
        except Exception:
            pass
        if spec.get("dataLabels"):
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showVal = True
    rows = int(math.ceil(chart.height / 0.53)) + 2
    return chart, rows


def block_chart(ctx: Ctx, ws: Any, b: Dict[str, Any], c0: int, r: int, span: int) -> int:
    ref = (b.get("data") or {}).get("table") or b.get("table") or b.get("source")
    if str(ref or "").lower() not in ctx.tables and not (not ref and len(ctx.tables) == 1):
        # the source table is defined on a sheet built later (e.g. a Dashboard placed first): reserve the space,
        # draw the chart once every sheet exists
        rows = int(math.ceil(float(b.get("height", 8.5)) / 0.53)) + 2
        ctx.deferred_charts.append((ws, b, c0, r))
        return r + rows
    reg = lookup_region(ctx, ref)
    chart, rows = build_chart(ctx, b, reg)
    place = str(b.get("place", "below"))
    if re.fullmatch(r"[A-Za-z]{1,3}\d+", place):
        ws.add_chart(chart, place.upper())
        pr, _ = parse_cell(place)
        return max(r, pr + rows)
    if place.lower() == "right":
        anchor_col = reg.last_col + 2
        top = max(1, reg.header_row - (1 if reg.header_row > 1 else 0))
        ws.add_chart(chart, f"{get_column_letter(anchor_col)}{top}")
        for cc in range(anchor_col, anchor_col + 9):  # keep room for the chart
            ctx.want_width(ws, cc, 9)
        return max(r, top + rows)
    ws.add_chart(chart, f"{get_column_letter(c0)}{r}")
    return r + rows


BLOCKS: Dict[str, Callable[..., int]] = {
    "title": block_title, "text": block_text, "notes": block_text, "kpis": block_kpis, "table": block_table,
    "summary": block_summary, "pivot": block_summary, "chart": block_chart,
}


def sheet_blocks(spec: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Accept explicit `blocks` or shorthand keys (title/kpis/headers+rows/charts/notes)."""
    blk = _maybe_json(spec.get("blocks"))
    if isinstance(blk, list) and blk:
        return [_maybe_json(x) for x in blk]
    blocks: List[Dict[str, Any]] = []
    if spec.get("title"):
        blocks.append({"type": "title", "text": spec["title"], "subtitle": spec.get("subtitle")})
    if spec.get("kpis"):
        blocks.append({"type": "kpis", "items": spec["kpis"]})
    if any(spec.get(k) for k in ("headers", "rows", "columns", "data", "values", "records")):
        blocks.append({"type": "table", **{k: spec[k] for k in (
            "headers", "header", "rows", "data", "values", "records", "columns", "totalRow", "excelTable", "banding",
            "filter", "freezeHeader", "conditionalFormats") if k in spec}, "name": spec.get("tableName", spec.get("name", "Data"))})
    for tb in spec.get("tables", []) or []:
        blocks.append({**tb, "type": "table"})
    for sm in spec.get("summaries", []) or []:
        blocks.append({**sm, "type": "summary"})
    tbl = next((b["name"] for b in blocks if b["type"] == "table"), None)
    for c in spec.get("charts", []) or []:
        c = dict(c)
        d = dict(c.get("data") or {})
        if tbl and not (d.get("table") or c.get("table") or c.get("source")):
            d["table"] = tbl  # shorthand charts default to this sheet's table
        if d:
            c["data"] = d
        if c.get("type") and str(c["type"]).lower() != "chart" and not c.get("chartType"):
            c["chartType"] = c["type"]  # {"type": "doughnut"} means chartType = doughnut
        c["type"] = "chart"
        blocks.append(c)
    if spec.get("notes"):
        blocks.append({"type": "text", "style": "note", "lines": spec["notes"] if isinstance(spec["notes"], list)
                       else [spec["notes"]]})
    return blocks


CHART_TYPE_NAMES = {"column", "bar", "stackedcolumn", "stackedbar", "percentstacked", "line", "area", "pie",
                    "doughnut", "donut", "scatter", "combo", "barchart", "columnchart", "linechart", "piechart"}
BLOCK_ALIASES = {"kpi": "kpis", "kpicards": "kpis", "cards": "kpis", "metrics": "kpis", "heading": "title",
                 "header": "title", "banner": "title", "paragraph": "text", "note": "text", "notes": "text",
                 "grid": "table", "data": "table", "datatable": "table", "pivot": "summary",
                 "pivottable": "summary", "chartblock": "chart", "graph": "chart", "gap": "spacer"}


def normalize_block(b: Any) -> Dict[str, Any]:
    """Be forgiving about how the agent names a block (e.g. {"type":"doughnut"} is a doughnut chart)."""
    if not isinstance(b, dict):
        raise ExcelError("Each block must be an object with a 'type'.", status=422, code="UNKNOWN_BLOCK")
    b = dict(b)
    t = re.sub(r"[^a-z]", "", str(b.get("type", "")).lower())
    if t in CHART_TYPE_NAMES:
        b.setdefault("chartType", t.replace("chart", ""))
        t = "chart"
    t = BLOCK_ALIASES.get(t, t)
    if not t:  # no type given: infer from content
        t = ("chart" if b.get("chartType") else "table" if (b.get("headers") or b.get("rows")) else
             "kpis" if (b.get("items") or b.get("kpis")) else "summary" if b.get("groupBy") else "text")
    b["type"] = t
    if t == "chart" and not b.get("chartType") and b.get("type") != "chart":
        b["chartType"] = "column"
    return b


def render_blocks(ctx: Ctx, ws: Any, blocks: List[Dict[str, Any]], origin: str = "B2") -> int:
    blocks = [normalize_block(b) for b in blocks]
    r, c0 = parse_cell(origin)
    span = 6
    for b in blocks:
        if b.get("type") == "table":
            try:
                span = max(span, len(normalize_table(b)[0]))
            except ExcelError:
                pass  # reported when the block is rendered
        elif b.get("type") == "summary":
            span = max(span, 1 + len(b.get("values") or [1]) + (1 if b.get("percentOfTotal") else 0))
        elif b.get("type") == "kpis":
            span = max(span, min(len(b.get("items") or b.get("kpis") or []), 4) * 3 - 1)
    span = min(span, 16)
    for b in blocks:
        t = str(b.get("type", "")).lower()
        if t == "spacer":
            r += int(b.get("rows", 1))
            continue
        fn = BLOCKS.get(t)
        if not fn:
            raise ExcelError(f"Unknown block type '{b.get('type')}'. Block types are: title, text, kpis, table, summary, "
                             "chart, spacer. For a chart use {\"type\":\"chart\",\"chartType\":\"doughnut\",...}.",
                             status=422, code="UNKNOWN_BLOCK")
        r = fn(ctx, ws, b, c0, r, span)
    return r


def apply_sheet_options(ctx: Ctx, ws: Any, spec: Dict[str, Any], index: int) -> None:
    ws.sheet_view.showGridLines = bool(spec.get("gridlines", False))
    ws.sheet_properties.tabColor = hex6(spec.get("tabColor"), ctx.theme.accents[index % len(ctx.theme.accents)])
    if spec.get("zoom"):
        ws.sheet_view.zoomScale = int(spec["zoom"])
    ws.page_setup.orientation = spec.get("orientation", "landscape")
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.left.text = "&A"


def build_sheet(ctx: Ctx, spec: Dict[str, Any], index: int) -> Any:
    name = safe_sheet_name(spec.get("name") or f"Sheet{index + 1}", ctx.wb.sheetnames)
    ws = ctx.wb.create_sheet(name)
    apply_sheet_options(ctx, ws, spec, index)
    origin = str(spec.get("origin", "B2"))
    if parse_cell(origin)[1] == 2:
        ws.column_dimensions["A"].width = 2
    render_blocks(ctx, ws, sheet_blocks(spec), origin)
    if spec.get("freeze"):
        ws.freeze_panes = str(spec["freeze"])
    ctx.note(f"Sheet '{name}' created")
    return ws


def build_workbook(spec: Dict[str, Any]) -> Ctx:
    sheets = _maybe_json(spec.get("sheets"))
    if isinstance(sheets, dict):
        sheets = [{"name": k, **(v if isinstance(v, dict) else {})} for k, v in sheets.items()]
    sheets = [_maybe_json(x) for x in sheets] if isinstance(sheets, list) else sheets
    if not isinstance(sheets, list) or not sheets:
        raise ExcelError("'sheets' must be a non-empty list. Each sheet needs a name and either 'blocks' or "
                         "shorthand fields (title, headers, rows, kpis, charts).", status=422, code="INVALID_SHEETS")
    if len(sheets) > 40:
        raise ExcelError("Too many sheets (max 40).", code="TOO_MANY_SHEETS")
    wb = Workbook()
    wb.remove(wb.active)
    ctx = Ctx(wb, Theme(spec.get("theme"), spec.get("currency")))
    for i, s in enumerate(sheets):
        if not isinstance(s, dict):
            raise ExcelError(f"Sheet {i + 1} must be an object.", status=422, code="INVALID_SHEETS")
        build_sheet(ctx, s, i)
    for nm in spec.get("namedRanges", []) or []:
        wb.defined_names[str(nm["name"])] = DefinedName(str(nm["name"]), attr_text=str(nm["ref"]))
    finalize_tokens(ctx)
    ctx.apply_widths()
    props = spec.get("properties") or {}
    wb.properties.title = str(props.get("title") or spec.get("title") or "Workbook")[:200]
    wb.properties.creator = str(props.get("author") or "Vkollab Technologies")[:100]
    wb.calculation.fullCalcOnLoad = True
    return ctx


# =========================================================================== #
# PART 2 - opening files, region detection, edit operations
# =========================================================================== #
def b64_decode(text: str) -> bytes:
    t = re.sub(r"^data:[^,]*,", "", str(text).strip())
    t = re.sub(r"\s+", "", t)
    t += "=" * (-len(t) % 4)
    try:
        return base64.b64decode(t.replace("-", "+").replace("_", "/"), validate=False)
    except (binascii.Error, ValueError):
        raise ExcelError("'fileBase64' is not valid base64.", code="INVALID_FILE")


class _NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *a, **k):  # follow redirects manually so every hop is checked
        return None


def _check_public_url(url: str) -> None:
    u = urlparse(url)
    if u.scheme not in ("http", "https") or not u.hostname:
        raise ExcelError("'fileUrl' must be an http(s) link.", code="INVALID_FILE_URL")
    try:
        infos = socket.getaddrinfo(u.hostname, u.port or (443 if u.scheme == "https" else 80))
    except socket.gaierror:
        raise ExcelError("The file link could not be resolved.", code="INVALID_FILE_URL")
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            raise ExcelError("That file link points to a private address and is not allowed.", code="INVALID_FILE_URL")


def fetch_url(url: str) -> bytes:
    """Download a publicly reachable file (e.g. an 'anyone with the link' direct download). SSRF-guarded."""
    opener = urllib.request.build_opener(_NoRedirect)
    for _ in range(4):
        _check_public_url(url)
        try:
            with opener.open(urllib.request.Request(url, headers={"User-Agent": "excel-service"}), timeout=25) as r:
                data = r.read(MAX_FILE_BYTES + 1)
                if len(data) > MAX_FILE_BYTES:
                    raise ExcelError(f"File is larger than {MAX_FILE_BYTES // 1048576} MB.", status=413, code="FILE_TOO_LARGE")
                return data
        except urllib.error.HTTPError as e:
            if e.code in (301, 302, 303, 307, 308) and e.headers.get("Location"):
                url = urllib.request.urljoin(url, e.headers["Location"])
                continue
            raise ExcelError(f"The file link returned HTTP {e.code}. Private SharePoint/OneDrive links need a "
                             "connector instead; pass the file as base64.", status=422, code="FILE_FETCH_FAILED")
        except (urllib.error.URLError, TimeoutError, OSError):
            raise ExcelError("The file could not be downloaded from that link.", status=422, code="FILE_FETCH_FAILED")
    raise ExcelError("Too many redirects.", status=422, code="FILE_FETCH_FAILED")


def load_source(payload: Dict[str, Any]) -> Tuple[bytes, str]:
    src = payload.get("source") if isinstance(payload.get("source"), dict) else payload
    name = str(src.get("fileName") or payload.get("sourceFileName") or "workbook.xlsx")
    if src.get("fileBase64"):
        data = b64_decode(src["fileBase64"])
    elif src.get("fileUrl"):
        data = fetch_url(str(src["fileUrl"]))
    else:
        raise ExcelError("Provide the workbook as 'source': {\"fileBase64\": \"...\", \"fileName\": \"x.xlsx\"} "
                         "or {\"fileUrl\": \"https://...\"}.", status=422, code="MISSING_FILE")
    if not data:
        raise ExcelError("The file is empty (or the base64 text was not valid).", status=400, code="INVALID_FILE")
    if len(data) > MAX_FILE_BYTES:
        raise ExcelError(f"File is larger than {MAX_FILE_BYTES // 1048576} MB.", status=413, code="FILE_TOO_LARGE")
    return data, name


NUM_RE = re.compile(r"^\(?\s*[-+]?\s*[$€£₹¥]?\s*[-+]?\s*(\d{1,3}(,\d{3})+|\d+)(\.\d+)?\s*%?\)?$")
DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d %b %Y", "%d-%b-%Y", "%b %d, %Y", "%B %d, %Y",
                "%d %B %Y", "%Y/%m/%d"]
US_DATE_FORMATS = ["%m/%d/%Y", "%m-%d-%Y"]


def parse_number(s: str) -> Optional[Tuple[Any, str]]:
    t = s.strip().replace("\xa0", " ")
    if not t or not NUM_RE.match(t) or re.match(r"^[-+]?0\d+$", t):
        return None
    digits = re.sub(r"[^\d.]", "", t)
    if not digits or digits.count(".") > 1 or len(digits.replace(".", "")) > 15:
        return None
    v = float(digits)
    neg = (t.startswith("(") and t.endswith(")")) or re.match(r"^\(?\s*-", t) is not None
    pct = t.endswith("%")
    if pct:
        v /= 100
    if neg:
        v = -v
    kind = "percent" if pct else "currency" if re.search(r"[$€£₹¥]", t) else "number"
    return (int(v) if float(v).is_integer() and "." not in digits and not pct else v), kind


def parse_date_text(s: str, day_first: bool = True) -> Optional[datetime]:
    t = s.strip()
    if not t or len(t) > 24 or not re.search(r"\d", t):
        return None
    for fmt in DATE_FORMATS + ([] if day_first else US_DATE_FORMATS):
        try:
            return datetime.strptime(t, fmt)
        except ValueError:
            continue
    if day_first:
        for fmt in US_DATE_FORMATS:  # only when day-first is impossible (e.g. 03/25/2026)
            try:
                return datetime.strptime(t, fmt)
            except ValueError:
                continue
    return None


def package_scan(data: bytes) -> Tuple[List[str], Dict[str, Any]]:
    """Find workbook features that openpyxl cannot preserve, so the agent can warn the user."""
    warnings: List[str] = []
    feats: Dict[str, Any] = {}
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return warnings, feats
    names = z.namelist()
    has = lambda prefix: any(n.startswith(prefix) for n in names)  # noqa: E731
    checks = [
        ("xl/slicers/", "Slicers will be lost when the file is saved."),
        ("xl/timelines/", "Timelines will be lost when the file is saved."),
        ("xl/threadedComments/", "Threaded comments will be converted or lost (classic notes are kept)."),
        ("xl/ctrlProps/", "Form controls (buttons, check boxes) will be lost."),
        ("xl/activeX/", "ActiveX controls will be lost."),
        ("xl/embeddings/", "Embedded objects may be lost."),
        ("xl/externalLinks/", "The workbook links to external workbooks; links are kept but not refreshed."),
        ("xl/queryTables/", "Query tables / Power Query connections may be lost."),
        ("xl/model/", "The Excel Data Model (Power Pivot) will be lost."),
    ]
    for prefix, msg in checks:
        if has(prefix):
            feats[prefix.strip("/").split("/")[-1]] = True
            warnings.append(msg)
    if has("xl/pivotTables/"):
        feats["pivotTables"] = True
        warnings.append("Pivot tables are kept but cannot be refreshed or edited here; refresh them in Excel.")
    if "xl/vbaProject.bin" in names:
        feats["macros"] = True
        warnings.append("The workbook contains macros: it will be saved as .xlsm to keep them (macros are not edited).")
    charts = [n for n in names if n.startswith("xl/charts/chart") and n.endswith(".xml")]
    if charts:
        feats["charts"] = len(charts)
        warnings.append(f"{len(charts)} existing chart(s) will be re-written; advanced chart formatting may change. "
                        "Check them after editing.")
    for n in names:
        if n.startswith("xl/worksheets/sheet") and n.endswith(".xml"):
            try:
                xml = z.read(n)[:3_000_000]
            except Exception:
                continue
            if b"sparklineGroup" in xml:
                feats["sparklines"] = True
            if b"x14:conditionalFormatting" in xml or b"x14:dataValidation" in xml:
                feats["x14Extensions"] = True
    if feats.get("sparklines"):
        warnings.append("Sparklines will be lost when the file is saved.")
    if feats.get("x14Extensions"):
        warnings.append("Some newer conditional formats / data validations (Excel 2010+ extensions) may be dropped.")
    for n in names:
        if n.startswith("xl/drawings/drawing") and n.endswith(".xml"):
            try:
                if b"<xdr:sp " in z.read(n)[:2_000_000]:
                    feats["shapes"] = True
            except Exception:
                pass
    if feats.get("shapes"):
        warnings.append("Shapes / text boxes on sheets may be lost (charts and pictures are kept).")
    return warnings, feats


def open_workbook(data: bytes, name: str, data_only: bool = False) -> Tuple[Workbook, str, str]:
    import warnings
    warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
    lower = name.lower()
    if data[:4] == b"PK\x03\x04":
        has_vba = False
        try:
            has_vba = "xl/vbaProject.bin" in zipfile.ZipFile(io.BytesIO(data)).namelist()
        except zipfile.BadZipFile:
            pass
        try:
            wb = load_workbook(io.BytesIO(data), data_only=data_only, keep_vba=has_vba or lower.endswith(".xlsm"))
        except Exception as exc:
            raise ExcelError("The file could not be opened as an Excel workbook (it may be corrupt).",
                             status=422, code="INVALID_WORKBOOK", details=str(exc)[:200])
        return wb, (".xlsm" if has_vba else ".xlsx"), "xlsx"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ExcelError("This looks like a legacy .xls file or a password-protected workbook. Save it as an "
                         "unprotected .xlsx in Excel and send it again.", status=422, code="UNSUPPORTED_FORMAT")
    try:  # CSV / TSV -> workbook
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    if "\x00" in text[:2000]:
        raise ExcelError("Unsupported file type. Send .xlsx, .xlsm or .csv.", status=422, code="UNSUPPORTED_FORMAT")
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    if len(rows) > MAX_ROWS_IN_SPEC:
        raise ExcelError(f"CSV has too many rows (limit {MAX_ROWS_IN_SPEC}).", status=413, code="FILE_TOO_LARGE")
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            if r > 1 and v.strip():
                num = parse_number(v)
                dt = None if num else parse_date_text(v)
                val: Any = num[0] if num else (dt if dt else clean_str(v))
            else:
                val = clean_str(v) if v.strip() else None
            ws.cell(r, c, val)
            if r > 1 and isinstance(val, datetime):
                ws.cell(r, c).number_format = "dd-mmm-yyyy"
    return wb, ".xlsx", "csv"


def detect_region(ws: Any) -> Optional[Region]:
    """Guess the main data table on a sheet: header row = first mostly-text row with 2+ cells."""
    if ws.max_row < 1 or ws.max_column < 1:
        return None
    max_c = min(ws.max_column, 200)
    head = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60), max_col=max_c, values_only=True))
    header_row = None
    for i, row in enumerate(head, 1):
        ne = [j for j, v in enumerate(row, 1) if v not in (None, "")]
        texty = [j for j in ne if isinstance(row[j - 1], str) and not is_formula(row[j - 1])]
        if len(ne) >= 2 and len(texty) >= 0.6 * len(ne):
            header_row = i
            break
    if header_row is None:
        for i, row in enumerate(head, 1):
            if any(v not in (None, "") for v in row):
                header_row = i
                break
    if header_row is None:
        return None
    ne = [j for j, v in enumerate(head[header_row - 1], 1) if v not in (None, "")]
    first_col, last_col = min(ne), max(ne)
    last_row, blank = header_row, 0
    for r, row in enumerate(ws.iter_rows(min_row=header_row + 1, max_row=min(ws.max_row, 200_000),
                                         min_col=first_col, max_col=last_col, values_only=True), header_row + 1):
        if any(v not in (None, "") for v in row):
            last_row, blank = r, 0
        else:
            blank += 1
            if blank >= 2:
                break
    total_row = None
    if last_row > header_row + 1:  # a trailing "Total" row is not data: keep it out of sums, charts, sorting
        vals = [ws.cell(last_row, c).value for c in range(first_col, last_col + 1)]
        first_text = next((v for v in vals if isinstance(v, str) and not is_formula(v)), "")
        agg = sum(1 for v in vals if is_formula(v) and re.match(r"=\s*(SUBTOTAL|SUM|AVERAGE)\(", v, re.I))
        if re.match(r"^\s*(grand\s+|sub\s*)?total", str(first_text), re.I) or agg >= max(1, len([v for v in vals if v not in (None, "")]) // 2):
            total_row, last_row = last_row, last_row - 1
    reg = region_from_bounds(ws, header_row, first_col, last_row, last_col, True)
    reg.total_row = total_row
    return reg


def region_from_bounds(ws: Any, r1: int, c1: int, r2: int, c2: int, has_header: bool = True) -> Region:
    headers: List[str] = []
    seen: Counter = Counter()
    for c in range(c1, c2 + 1):
        v = ws.cell(r1, c).value if has_header else None
        h = clean_str(str(v)).strip() if v not in (None, "") else (f"Column {get_column_letter(c)}")
        seen[h.lower()] += 1
        headers.append(h if seen[h.lower()] == 1 else f"{h}_{seen[h.lower()]}")
    first = r1 + 1 if has_header else r1
    fmts = {}
    if r2 >= first:
        for i, h in enumerate(headers):
            nf = ws.cell(first, c1 + i).number_format
            if nf and nf != "General":
                fmts[h] = nf
    return Region(ws.title, r1 if has_header else max(r1 - 1, 1), c1, c2, first, r2, headers, fmts)


def op_region(ctx: Ctx, p: Dict[str, Any], ws: Any) -> Region:
    t = str(p.get("table") or "").lower()
    if t and t in ctx.tables:
        return ctx.tables[t]
    if p.get("range"):
        r1, c1, r2, c2 = parse_range(ws, p["range"])
        return region_from_bounds(ws, r1, c1, r2, c2, bool(p.get("hasHeader", True)))
    reg = detect_region(ws)
    if not reg:
        raise ExcelError(f"Sheet '{ws.title}' has no data to work on.", status=422, code="EMPTY_SHEET")
    return reg


# ---- style helpers ----------------------------------------------------------- #
def color_of(theme: Theme, c: Any, default: Optional[str] = None) -> Optional[str]:
    if c is None:
        return default
    key = str(c).strip().lower()
    if key in ("primary", "dark", "light", "band", "border", "text", "muted"):
        return getattr(theme, key)
    m = re.fullmatch(r"accent([1-6])", key)
    if m:
        return theme.accents[int(m.group(1)) - 1]
    if key in COLOR_ALIASES:
        return COLOR_ALIASES[key][0]
    return hex6(c, default)


def apply_style(cell: Any, p: Dict[str, Any], theme: Theme) -> None:
    f = p.get("font")
    if isinstance(f, dict):
        cur = cell.font
        cell.font = Font(name=f.get("name", cur.name), size=f.get("size", cur.sz), bold=f.get("bold", cur.b),
                         italic=f.get("italic", cur.i), underline=f.get("underline", cur.u),
                         strike=f.get("strike", cur.strike),
                         color=color_of(theme, f.get("color")) or cur.color)
    fl = p.get("fill")
    if fl is not None:
        color = fl.get("color") if isinstance(fl, dict) else fl
        if str(color).lower() in ("none", "clear", ""):
            cell.fill = PatternFill(fill_type=None)
        else:
            cc = color_of(theme, color)
            if cc:
                cell.fill = fill_of(cc)
    if p.get("numberFormat"):
        cell.number_format = number_format(p["numberFormat"], theme) or "General"
    a = p.get("alignment")
    if isinstance(a, dict):
        cur = cell.alignment
        cell.alignment = Alignment(horizontal=a.get("horizontal", cur.horizontal), vertical=a.get("vertical", cur.vertical),
                                   wrap_text=a.get("wrap", cur.wrap_text), indent=a.get("indent", cur.indent),
                                   text_rotation=a.get("rotation", cur.text_rotation))


def apply_border(ws: Any, rng: str, b: Dict[str, Any], theme: Theme) -> None:
    r1, c1, r2, c2 = parse_range(ws, rng)
    style = b.get("style", "thin")
    s = side_of(color_of(theme, b.get("color"), theme.border) or theme.border, style)
    sides = b.get("sides", "all")
    sides = [sides] if isinstance(sides, str) else sides
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell, cur = ws.cell(r, c), ws.cell(r, c).border
            allb, outer = "all" in sides, "outer" in sides
            left = s if allb or "left" in sides or (outer and c == c1) else cur.left
            right = s if allb or "right" in sides or (outer and c == c2) else cur.right
            top = s if allb or "top" in sides or (outer and r == r1) else cur.top
            bottom = s if allb or "bottom" in sides or (outer and r == r2) else cur.bottom
            cell.border = Border(left=left, right=right, top=top, bottom=bottom)


def autofit_columns(ws: Any, c1: int, c2: int, r1: int = 1, r2: Optional[int] = None) -> None:
    r2 = min(r2 or ws.max_row, r1 + 2000)
    for c in range(c1, c2 + 1):
        best = 0.0
        for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c, max_col=c):
            cell = row[0]
            v = cell.value
            if v is None:
                continue
            if isinstance(v, (datetime, date)):
                n = 12
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                n = len(f"{v:,.2f}") + (2 if "%" in cell.number_format or "$" in cell.number_format else 0)
            elif is_formula(v):
                n = 12
            else:
                n = max(len(x) for x in str(v).split("\n"))
                if cell.alignment.wrap_text:
                    n = min(n, 40)
            best = max(best, n * (1.15 if cell.font and cell.font.b else 1.0))
        ws.column_dimensions[get_column_letter(c)].width = max(8, min(best + 2.5, 60))


# ---- operations --------------------------------------------------------------- #
OPS: Dict[str, Callable[[Ctx, Dict[str, Any]], None]] = {}
OP_DOCS: Dict[str, str] = {}


def op(*names: str, doc: str = ""):
    def deco(fn):
        for n in names:
            OPS[re.sub(r"[^a-z]", "", n.lower())] = fn
        OP_DOCS[names[0]] = doc
        return fn
    return deco



# ---- keep formulas correct when rows/columns are inserted or deleted --------------- #
from bisect import bisect_left, bisect_right  # noqa: E402
from openpyxl.formula import Tokenizer  # noqa: E402

_CELL = re.compile(r"(\$?)([A-Za-z]{1,3})(\$?)(\d+)")
_COLRANGE = re.compile(r"(\$?)([A-Za-z]{1,3}):(\$?)([A-Za-z]{1,3})")
_ROWRANGE = re.compile(r"(\$?)(\d+):(\$?)(\d+)")


def _map_index(kind: str, a: int, b: int, positions: List[int], count: int) -> Optional[Tuple[int, int]]:
    """New (start, end) of the span a..b after deleting `positions` (kind='del') or inserting
    `count` rows/cols at positions[0] (kind='ins'). None => the whole span was deleted (#REF!)."""
    if kind == "ins":
        at = positions[0]
        return (a + count if a >= at else a, b + count if b >= at else b)
    na, nb = a - bisect_left(positions, a), b - bisect_right(positions, b)
    return (na, nb) if nb >= na else None


def _shift_ref(ref: str, axis: str, kind: str, positions: List[int], count: int) -> str:
    def col_i(x: str) -> int:
        return column_index_from_string(x.upper())

    parts = ref.split(":")
    if len(parts) == 2 and _COLRANGE.fullmatch(ref):
        if axis != "col":
            return ref
        m = _COLRANGE.fullmatch(ref)
        res = _map_index(kind, col_i(m.group(2)), col_i(m.group(4)), positions, count)
        return "#REF!" if not res else f"{m.group(1)}{get_column_letter(res[0])}:{m.group(3)}{get_column_letter(res[1])}"
    if len(parts) == 2 and _ROWRANGE.fullmatch(ref):
        if axis != "row":
            return ref
        m = _ROWRANGE.fullmatch(ref)
        res = _map_index(kind, int(m.group(2)), int(m.group(4)), positions, count)
        return "#REF!" if not res else f"{m.group(1)}{res[0]}:{m.group(3)}{res[1]}"
    cells = [_CELL.fullmatch(x) for x in parts]
    if not cells or any(c is None for c in cells) or len(cells) > 2:
        return ref  # named range, function result, etc.
    vals = [(col_i(c.group(2)) if axis == "col" else int(c.group(4))) for c in cells]
    res = _map_index(kind, vals[0], vals[-1], positions, count)
    if not res:
        return "#REF!"
    out = []
    for c, new in zip(cells, [res[0], res[1]][: len(cells)] if len(cells) == 2 else [res[0]]):
        out.append(f"{c.group(1)}{get_column_letter(new)}{c.group(3)}{c.group(4)}" if axis == "col"
                   else f"{c.group(1)}{c.group(2)}{c.group(3)}{new}")
    return ":".join(out)


def shift_workbook_formulas(wb: Workbook, target: Any, axis: str, kind: str, positions: List[int], count: int = 1) -> int:
    """Rewrite formula references (any sheet) after rows/cols were/are about to be inserted or deleted in `target`."""
    changed = 0
    positions = sorted(positions)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type != "f" or not isinstance(c.value, str):
                    continue
                try:
                    tok = Tokenizer(c.value)
                    out = []
                    for t in tok.items:
                        v = t.value
                        if t.type == "OPERAND" and t.subtype == "RANGE":
                            if "!" in v:
                                sheet_part, ref = v.rsplit("!", 1)
                                sheet = sheet_part.strip("'").replace("''", "'")
                                prefix = sheet_part + "!"
                            else:
                                sheet, ref, prefix = ws.title, v, ""
                            if sheet == target.title:
                                new = _shift_ref(ref, axis, kind, positions, count)
                                v = new if new == "#REF!" else prefix + new
                        out.append(v)
                    new_f = "=" + "".join(out)
                except Exception:
                    continue
                if new_f != c.value:
                    c.value = new_f
                    changed += 1
    return changed


def _formula_warn(ctx: Ctx, ws: Any) -> None:
    """Formulas are rewritten automatically; these other objects are not shifted."""
    things = []
    if ws.merged_cells.ranges:
        things.append("merged ranges")
    if any(True for _ in ws.conditional_formatting):
        things.append("conditional formats")
    if ws.data_validations and ws.data_validations.dataValidation:
        things.append("data validations")
    if getattr(ws, "_charts", None):
        things.append("chart ranges")
    if things:
        ctx.warn(f"Rows/columns changed on '{ws.title}': {', '.join(things)} are not shifted automatically; review them. "
                 "Formula references were updated.")


def _put(ws: Any, row: int, col: int, value: Any) -> None:
    try:
        ws.cell(row, col).value = value
    except AttributeError:
        raise ExcelError(f"{get_column_letter(col)}{row} is inside a merged range; write to the top-left cell of "
                         "the merge, or unmerge it first.", status=422, code="MERGED_CELL")


@op("set_cells", doc="{sheet, cells:{'A1': value|'=formula'}} or {sheet, at:'B2', values:[[..],[..]]}")
def op_set_cells(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    n = 0
    if isinstance(p.get("cells"), dict):
        for ref, v in p["cells"].items():
            row, col = parse_cell(ref)
            _put(ws, row, col, cell_value(v, row))
            n += 1
    elif p.get("values"):
        r0, c0 = parse_cell(p.get("at", "A1"))
        for i, rowv in enumerate(p["values"]):
            for j, v in enumerate(rowv if isinstance(rowv, list) else [rowv]):
                _put(ws, r0 + i, c0 + j, cell_value(v, r0 + i))
                n += 1
    else:
        raise ExcelError("set_cells needs 'cells' or 'values'.", status=422, code="INVALID_OPERATION")
    ctx.note(f"Set {n} cell(s) on '{ws.title}'")


@op("clear_range", doc="{sheet, range, what: 'contents'|'formats'|'all'}")
def op_clear_range(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    what = str(p.get("what", "contents")).lower()
    r1, c1, r2, c2 = parse_range(ws, p["range"])
    for mr in list(ws.merged_cells.ranges):  # merged cells are read-only: unmerge anything we overlap
        if not (mr.max_row < r1 or mr.min_row > r2 or mr.max_col < c1 or mr.min_col > c2):
            ws.unmerge_cells(str(mr))
            ctx.note(f"Unmerged {mr} so it could be cleared")
    for c in range_cells(ws, p["range"]):
        if what in ("contents", "all"):
            c.value = None
        if what in ("formats", "all"):
            c.style = "Normal"
    ctx.note(f"Cleared {what} in {p['range']} on '{ws.title}'")


@op("format_range", "format", doc="{sheet, range, font:{bold,italic,size,color,name}, fill:'primary'|'#hex', numberFormat, alignment:{horizontal,vertical,wrap,indent}, border:{style,color,sides}}")
def op_format_range(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    for c in range_cells(ws, p["range"]):
        apply_style(c, p, ctx.theme)
    if isinstance(p.get("border"), dict):
        apply_border(ws, p["range"], p["border"], ctx.theme)
    ctx.note(f"Formatted {p['range']} on '{ws.title}'")


@op("set_column_width", "column_width", doc="{sheet, columns:{'A':20,'C':12}}")
def op_set_column_width(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    for col, w in (p.get("columns") or {}).items():
        ws.column_dimensions[str(col).upper()].width = float(w)
    ctx.note(f"Set widths on '{ws.title}'")


@op("set_row_height", "row_height", doc="{sheet, rows:{'1':30}}")
def op_set_row_height(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    for row, h in (p.get("rows") or {}).items():
        ws.row_dimensions[int(row)].height = float(h)
    ctx.note(f"Set row heights on '{ws.title}'")


@op("autofit_columns", "autofit", doc="{sheet, range?}")
def op_autofit(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    if p.get("range"):
        r1, c1, r2, c2 = parse_range(ws, p["range"])
    else:
        r1, c1, r2, c2 = 1, 1, ws.max_row, ws.max_column
    autofit_columns(ws, c1, c2, r1, r2)
    ctx.note(f"Auto-fitted columns on '{ws.title}'")


@op("merge_cells", "merge", doc="{sheet, range}")
def op_merge(ctx, p):
    ctx.sheet(p.get("sheet")).merge_cells(p["range"])
    ctx.note(f"Merged {p['range']}")


@op("unmerge_cells", "unmerge", doc="{sheet, range}")
def op_unmerge(ctx, p):
    ctx.sheet(p.get("sheet")).unmerge_cells(p["range"])
    ctx.note(f"Unmerged {p['range']}")


@op("freeze_panes", "freeze", doc="{sheet, cell:'A2'|null}")
def op_freeze(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    ws.freeze_panes = p.get("cell") or None
    ctx.note(f"Freeze panes on '{ws.title}' at {p.get('cell') or 'none'}")


@op("add_filter", "autofilter", doc="{sheet, range?}")
def op_filter(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    ws.auto_filter.ref = p.get("range") or op_region(ctx, p, ws).data_ref
    ctx.note(f"Filter added on '{ws.title}' ({ws.auto_filter.ref})")


def _sort_key(v: Any):
    if v is None or v == "":
        return (2, 0, "")
    if isinstance(v, bool):
        return (1, 0, str(v))
    if isinstance(v, (int, float)):
        return (0, float(v), "")
    if isinstance(v, (datetime, date)):
        return (0, datetime(v.year, v.month, v.day).timestamp() if not isinstance(v, datetime) else v.timestamp(), "")
    return (1, 0, str(v).lower())


@op("sort_range", "sort", doc="{sheet, range?, hasHeader?, by:[{column:'Revenue', order:'desc'}]}")
def op_sort(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    reg = op_region(ctx, p, ws)
    by = p.get("by") or ([{"column": p["column"], "order": p.get("order", "asc")}] if p.get("column") else None)
    if not by:
        raise ExcelError("sort_range needs 'by': [{column, order}].", status=422, code="INVALID_OPERATION")
    if reg.last_row <= reg.first_row:
        return
    rows = [[ws.cell(r, c).value for c in range(reg.first_col, reg.last_col + 1)] for r in range(reg.first_row, reg.last_row + 1)]
    if any(is_formula(v) for row in rows for v in row) and not p.get("allowFormulas"):
        raise ExcelError("The range contains formulas; sorting would break their references. Sort a values-only "
                         "range or pass allowFormulas=true.", status=422, code="SORT_FORMULAS")
    for spec in reversed(by):  # stable multi-key sort
        ci = reg.col(spec["column"]) - reg.first_col
        desc = str(spec.get("order", "asc")).lower().startswith("d")
        present = [r for r in rows if r[ci] not in (None, "")]
        blanks = [r for r in rows if r[ci] in (None, "")]
        present.sort(key=lambda r: _sort_key(r[ci]), reverse=desc)
        rows = present + blanks
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            ws.cell(reg.first_row + i, reg.first_col + j).value = v
    ctx.note(f"Sorted '{ws.title}' by " + ", ".join(f"{s['column']} {s.get('order', 'asc')}" for s in by))


@op("find_replace", doc="{sheet?, find, replace, matchCase?, wholeCell?, regex?, inFormulas?}")
def op_find_replace(ctx, p):
    sheets = [ctx.sheet(p["sheet"])] if p.get("sheet") else ctx.wb.worksheets
    find, rep = str(p.get("find", "")), str(p.get("replace", ""))
    if not find:
        raise ExcelError("find_replace needs 'find'.", status=422, code="INVALID_OPERATION")
    flags = 0 if p.get("matchCase") else re.I
    pat = re.compile(find if p.get("regex") else re.escape(find), flags)
    n = 0
    for ws in sheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str) or (is_formula(v) and not p.get("inFormulas")):
                    continue
                if p.get("wholeCell"):
                    if pat.fullmatch(v.strip()):
                        c.value, n = rep, n + 1
                else:
                    nv, k = pat.subn(rep, v)
                    if k:
                        c.value, n = nv, n + k
    ctx.note(f"Replaced {n} occurrence(s) of '{find}'")


AUTO_FORMATS = ("General", "yyyy-mm-dd", "yyyy-mm-dd h:mm:ss", "mm-dd-yy")  # formats Excel/openpyxl assign by default


def convert_types(ws: Any, reg: Region, theme: Theme, columns: Optional[List[int]] = None,
                  day_first: bool = True) -> Counter:
    """Turn text that is really a number/date into real numbers/dates."""
    counts: Counter = Counter()
    cols = columns or list(range(reg.first_col, reg.last_col + 1))
    for c in cols:
        cells = [ws.cell(r, c) for r in range(reg.first_row, reg.last_row + 1)]
        texts = [x for x in cells if isinstance(x.value, str) and not is_formula(x.value) and x.value.strip()]
        if not texts:
            continue
        # only convert a column when (nearly) all its text cells convert: avoids touching IDs / names
        nums = [parse_number(x.value) for x in texts]
        if all(nums) and len(texts) >= 1:
            for x, nres in zip(texts, nums):
                x.value = nres[0]
                if x.number_format == "General":
                    x.number_format = {"percent": "0.0%", "currency": f'"{theme.symbol}"#,##0.00'}.get(
                        nres[1], "#,##0.00" if isinstance(nres[0], float) else "General")
                counts["text_to_numbers"] += 1
            continue
        dts = [parse_date_text(x.value, day_first) for x in texts]
        if all(dts):
            for x, d in zip(texts, dts):
                x.value = d
                if x.number_format in AUTO_FORMATS:
                    x.number_format = "dd-mmm-yyyy"
                counts["text_to_dates"] += 1
    return counts


@op("clean_data", "clean", doc="{sheet, range?, actions:['trim','collapse_spaces','text_to_numbers','text_to_dates','remove_blank_rows','dedupe','proper_case','upper_case','lower_case','fill_blanks'], columns?, keys?, fillValue?, dayFirst?}")
def op_clean(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    reg = op_region(ctx, p, ws)
    actions = {re.sub(r"[^a-z_]", "", str(a).lower()) for a in (
        p.get("actions") or ["trim", "collapse_spaces", "text_to_numbers", "text_to_dates", "remove_blank_rows"])}
    cols = [reg.col(c) for c in p["columns"]] if p.get("columns") else list(range(reg.first_col, reg.last_col + 1))
    counts: Counter = Counter()
    for r in range(reg.first_row, reg.last_row + 1):
        for c in cols:
            cell = ws.cell(r, c)
            v = cell.value
            if isinstance(v, str) and not is_formula(v):
                nv = v.replace("\xa0", " ")
                if "trim" in actions:
                    nv = nv.strip()
                if "collapse_spaces" in actions:
                    nv = re.sub(r"[ \t]+", " ", nv)
                if "proper_case" in actions:
                    nv = nv.title()
                elif "upper_case" in actions:
                    nv = nv.upper()
                elif "lower_case" in actions:
                    nv = nv.lower()
                if nv != v:
                    cell.value = nv
                    counts["text_cleaned"] += 1
            elif v in (None, "") and "fill_blanks" in actions and p.get("fillValue") is not None:
                cell.value = coerce(p["fillValue"])
                counts["blanks_filled"] += 1
    if actions & {"text_to_numbers", "text_to_dates"}:
        counts.update(convert_types(ws, reg, ctx.theme, cols if p.get("columns") else None, p.get("dayFirst", True)))
    if ws.merged_cells.ranges and actions & {"remove_blank_rows", "dedupe"}:
        ctx.warn(f"'{ws.title}' has merged cells; removing rows may misalign them.")
    to_delete: List[int] = []
    if "remove_blank_rows" in actions:
        for r in range(reg.first_row, reg.last_row + 1):
            if all(ws.cell(r, c).value in (None, "") for c in range(reg.first_col, reg.last_col + 1)):
                to_delete.append(r)
        counts["blank_rows_removed"] += len(to_delete)
    if "dedupe" in actions:
        keys = [reg.col(k) for k in p["keys"]] if p.get("keys") else list(range(reg.first_col, reg.last_col + 1))
        seen = set()
        for r in range(reg.first_row, reg.last_row + 1):
            if r in to_delete:
                continue
            key = tuple(str(ws.cell(r, c).value).strip().lower() for c in keys)
            if key in seen:
                to_delete.append(r)
                counts["duplicates_removed"] += 1
            seen.add(key)
    if to_delete:
        shift_workbook_formulas(ctx.wb, ws, "row", "del", sorted(set(to_delete)))
        _formula_warn(ctx, ws)
    for r in sorted(set(to_delete), reverse=True):
        ws.delete_rows(r)
    ctx.note(f"Cleaned '{ws.title}': " + (", ".join(f"{k.replace('_', ' ')} {v}" for k, v in counts.items()) or "nothing to change"))


@op("add_formula_column", "add_column", doc="{sheet, header, formula:'=[@Units]*[@Price]', after?, format?, range?}")
def op_add_formula_column(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    reg = op_region(ctx, p, ws)
    header = str(p.get("header") or "New column")
    letters = {h.lower(): get_column_letter(reg.first_col + i) for i, h in enumerate(reg.headers)}
    if p.get("after"):
        col = reg.col(p["after"]) + 1
        shift_workbook_formulas(ctx.wb, ws, "col", "ins", [col], 1)
        ws.insert_cols(col)
        _formula_warn(ctx, ws)
        letters = {h.lower(): get_column_letter(reg.first_col + i + (1 if reg.first_col + i >= col else 0))
                   for i, h in enumerate(reg.headers)}
    else:
        col = reg.last_col + 1
    ref_col = max(col - 1, reg.first_col)
    hc = ws.cell(reg.header_row, col, header)
    hc._style = copy_style(ws.cell(reg.header_row, ref_col))
    fmt = number_format(p.get("format"), ctx.theme)
    for r in range(reg.first_row, reg.last_row + 1):
        c = ws.cell(r, col, resolve_tokens(ctx, cell_value(p.get("formula", ""), r), letters, r))
        c._style = copy_style(ws.cell(r, ref_col))
        if fmt:
            c.number_format = fmt
    autofit_columns(ws, col, col, reg.header_row, reg.last_row)
    if ws.auto_filter.ref:
        c1, r1, c2, r2 = range_boundaries(ws.auto_filter.ref)
        if r1 == reg.header_row and c1 <= col <= c2 + 1:
            ws.auto_filter.ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(max(c2, col))}{r2}"
    ctx.note(f"Added column '{header}' to '{ws.title}'")


def copy_style(cell: Any) -> Any:
    from copy import copy
    return copy(cell._style)


@op("add_totals_row", "add_totals", doc="{sheet, range?, columns:{'Revenue':'sum','Units':'avg'}}")
def op_add_totals(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    reg = op_region(ctx, p, ws)
    codes = {"sum": 109, "avg": 101, "average": 101, "count": 102, "counta": 103, "max": 104, "min": 105}
    row = reg.last_row + 1
    th = ctx.theme
    for c in range(reg.first_col, reg.last_col + 1):
        cell = ws.cell(row, c)
        cell.font = Font(name=th.font, bold=True, color=th.dark)
        cell.fill = fill_of(th.light)
        cell.border = Border(top=side_of(th.dark, "medium"))
    ws.cell(row, reg.first_col, "Total")
    for header, agg in (p.get("columns") or {}).items():
        col = reg.col(header)
        L = get_column_letter(col)
        a = str(agg).lower()
        if a not in codes:
            raise ExcelError(f"Unknown total '{agg}'. Use sum, avg, count, counta, min or max.", status=422, code="INVALID_OPERATION")
        cell = ws.cell(row, col, f"=SUBTOTAL({codes[a]},{L}{reg.first_row}:{L}{reg.last_row})")
        cell.number_format = ws.cell(reg.last_row, col).number_format
        cell.alignment = Alignment(horizontal="right")
    ctx.note(f"Added totals row to '{ws.title}' (row {row})")


@op("enhance_sheet", "enhance", "format_table", "style_table", doc="{sheet, range?, options:{fixTypes, header, numberFormats, banding, borders, widths, freeze, filter, hideGridlines, unifyFont}} (all default true except unifyFont)")
def op_enhance(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    reg = op_region(ctx, p, ws)
    th = ctx.theme
    o = p.get("options") if isinstance(p.get("options"), dict) else p

    def flag(k: str, d: bool = True) -> bool:
        return bool(o.get(k, d))

    done = []
    if flag("fixTypes"):
        n = convert_types(ws, reg, th, day_first=o.get("dayFirst", True))
        if n:
            done.append("fixed " + ", ".join(f"{v} {k.replace('_', ' ')}" for k, v in n.items()))
    if flag("header"):
        for c in range(reg.first_col, reg.last_col + 1):
            cell = ws.cell(reg.header_row, c)
            cell.font = Font(name=th.font, size=11, bold=True, color="FFFFFF")
            cell.fill = fill_of(th.dark)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=side_of(th.primary, "medium"))
        ws.row_dimensions[reg.header_row].height = max(ws.row_dimensions[reg.header_row].height or 15, 28)
        done.append("styled header")
    rows = range(reg.first_row, reg.last_row + 1)
    for ci, h in enumerate(reg.headers):
        col = reg.first_col + ci
        vals = [ws.cell(r, col).value for r in list(rows)[:500]]
        if flag("numberFormats"):
            key = infer_format(h, vals)
            fmt = number_format(key, th) if key else None
            if fmt:
                for r in rows:
                    cell = ws.cell(r, col)
                    if cell.value in (None, ""):
                        continue
                    if cell.number_format == "General" or (fmt == "dd-mmm-yyyy" and cell.number_format in AUTO_FORMATS):
                        cell.number_format = fmt
        for r in rows:
            cell = ws.cell(r, col)
            if cell.alignment.horizontal is None and cell.value is not None:
                if isinstance(cell.value, (datetime, date)):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
    if flag("numberFormats"):
        done.append("applied number formats")
    for r in rows:
        for c in range(reg.first_col, reg.last_col + 1):
            cell = ws.cell(r, c)
            if flag("banding") and (r - reg.first_row) % 2 == 1 and cell.fill.fill_type is None:
                cell.fill = fill_of(th.band)
            if flag("borders") and not (cell.border.bottom and cell.border.bottom.style):
                cell.border = Border(bottom=side_of(th.border))
            if flag("unifyFont", False):
                f = cell.font
                cell.font = Font(name=th.font, size=f.sz, bold=f.b, italic=f.i, color=f.color)
    if flag("banding") or flag("borders"):
        done.append("banding and borders")
    if flag("widths"):
        autofit_columns(ws, reg.first_col, reg.last_col, reg.header_row, reg.last_row)
        done.append("fitted column widths")
    if flag("freeze") and not ws.freeze_panes and reg.header_row <= 12:
        ws.freeze_panes = ws.cell(reg.header_row + 1, 1)
        done.append("froze header")
    if flag("filter") and not ws.auto_filter.ref and not ws.merged_cells.ranges:
        ws.auto_filter.ref = reg.data_ref
        done.append("added filter")
    if flag("hideGridlines"):
        ws.sheet_view.showGridLines = False
    ctx.note(f"Enhanced '{ws.title}' ({reg.data_ref}): " + ", ".join(done))


@op("conditional_format", "add_conditional_format", doc="{sheet, range|table+column, rule:{type:'colorScale'|'dataBar'|'iconSet'|'cell'|'text'|'formula'|'topN'|'bottomN'|'duplicates'|'blanks'|'status', ...}}")
def op_cf(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    rule = p.get("rule") or p
    rng = p.get("range")
    if not rng and p.get("column"):
        reg = op_region(ctx, p, ws)
        L = get_column_letter(reg.col(p["column"]))
        rng = f"{L}{reg.first_row}:{L}{reg.last_row}"
    if not rng:
        raise ExcelError("conditional_format needs 'range' (or 'column').", status=422, code="INVALID_OPERATION")
    add_cf(ws, rng, rule, ctx.theme)
    ctx.note(f"Conditional format ({rule.get('type')}) on {rng}")


@op("data_validation", "add_validation", doc="{sheet, range|column, rule:{type:'list'|'whole'|'decimal'|'date'|'textLength'|'custom', values|source|min|max|operator|formula, error, prompt}}")
def op_dv(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    rule = p.get("rule") or p
    rng = p.get("range")
    if not rng and p.get("column"):
        reg = op_region(ctx, p, ws)
        L = get_column_letter(reg.col(p["column"]))
        rng = f"{L}{reg.first_row}:{L}{reg.last_row + int(p.get('extraRows', 200))}"
    if not rng:
        raise ExcelError("data_validation needs 'range' (or 'column').", status=422, code="INVALID_OPERATION")
    add_validation(ctx, ws, rng, rule)
    ctx.note(f"Data validation ({rule.get('type', 'list')}) on {rng}")


@op("add_chart", doc="{sheet (source), range?|table?, categories, series:[..], chartType, title, place:'below'|'right'|'H2', targetSheet?, width, height, dataLabels, legend, xTitle, yTitle}")
def op_add_chart(ctx, p):
    src = ctx.sheet(p.get("sheet"))
    reg = op_region(ctx, p, src)
    target = ctx.sheet(p["targetSheet"]) if p.get("targetSheet") else src
    chart, rows = build_chart(ctx, p, reg)
    place = str(p.get("place", "right"))
    if re.fullmatch(r"[A-Za-z]{1,3}\d+", place):
        anchor = place.upper()
    elif place.lower() == "below":
        anchor = f"{get_column_letter(reg.first_col)}{reg.last_row + 3}"
    else:
        anchor = f"{get_column_letter(reg.last_col + 2)}{max(reg.header_row, 1)}"
    target.add_chart(chart, anchor)
    ctx.note(f"Added {p.get('chartType', p.get('type', 'column'))} chart on '{target.title}' at {anchor}")


@op("add_blocks", doc="{sheet, at?:'B2', blocks:[title|text|kpis|table|summary|chart...]} - render formatted blocks into an existing sheet (below existing data by default)")
def op_add_blocks(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    origin = p.get("at") or f"B{ws.max_row + 3}"
    render_blocks(ctx, ws, list(p.get("blocks") or []), origin)
    finalize_tokens(ctx)
    ctx.note(f"Added blocks to '{ws.title}' at {origin}")


@op("add_sheet", doc="{name, blocks:[..] | shorthand fields (title, headers, rows, kpis, charts), position?}")
def op_add_sheet(ctx, p):
    ws = build_sheet(ctx, p, len(ctx.wb.sheetnames))
    if p.get("position") is not None:
        ctx.wb.move_sheet(ws, offset=int(p["position"]) - ctx.wb.sheetnames.index(ws.title))
    finalize_tokens(ctx)


@op("rename_sheet", doc="{sheet, newName}")
def op_rename_sheet(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    old = ws.title
    new = safe_sheet_name(p.get("newName"), [s for s in ctx.wb.sheetnames if s != old])
    ws.title = new
    pat = re.compile(r"(?<![A-Za-z0-9_])(?:'" + re.escape(old.replace("'", "''")) + r"'|" + re.escape(old) + r")!")
    n = 0
    for sh in ctx.wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if is_formula(c.value) and pat.search(c.value):
                    c.value, k = pat.subn(q(new) + "!", c.value)
                    n += 1
    ctx.note(f"Renamed sheet '{old}' to '{new}'" + (f" (updated {n} formula(s))" if n else ""))


@op("delete_sheet", doc="{sheet}")
def op_delete_sheet(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    if len(ctx.wb.sheetnames) == 1:
        raise ExcelError("A workbook must keep at least one sheet.", status=422, code="INVALID_OPERATION")
    ctx.warn(f"Deleted sheet '{ws.title}'. Formulas elsewhere that referenced it will show #REF!.")
    ctx.wb.remove(ws)
    ctx.note(f"Deleted sheet '{ws.title}'")


@op("copy_sheet", "duplicate_sheet", doc="{sheet, newName?}")
def op_copy_sheet(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    cp = ctx.wb.copy_worksheet(ws)
    cp.title = safe_sheet_name(p.get("newName") or f"{ws.title} copy", [s for s in ctx.wb.sheetnames if s != cp.title])
    ctx.note(f"Copied '{ws.title}' to '{cp.title}'")


@op("move_sheet", doc="{sheet, position (0-based)}")
def op_move_sheet(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    ctx.wb.move_sheet(ws, offset=int(p.get("position", 0)) - ctx.wb.sheetnames.index(ws.title))
    ctx.note(f"Moved '{ws.title}' to position {p.get('position', 0)}")


@op("sheet_properties", "set_sheet", doc="{sheet, tabColor?, hidden?, gridlines?, zoom?}")
def op_sheet_props(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    if p.get("tabColor"):
        ws.sheet_properties.tabColor = color_of(ctx.theme, p["tabColor"], ctx.theme.primary)
    if "hidden" in p:
        ws.sheet_state = "hidden" if p["hidden"] else "visible"
    if "gridlines" in p:
        ws.sheet_view.showGridLines = bool(p["gridlines"])
    if p.get("zoom"):
        ws.sheet_view.zoomScale = int(p["zoom"])
    ctx.note(f"Updated properties of '{ws.title}'")


def _shift(ctx: Ctx, p: Dict[str, Any], fn: str, label: str) -> None:
    ws = ctx.sheet(p.get("sheet"))
    idx = p.get("at") or p.get("index")
    if idx is None:
        raise ExcelError(f"{label} needs 'at' (row number or column letter).", status=422, code="INVALID_OPERATION")
    count = int(p.get("count", 1))
    if "col" in fn:
        idx = column_index_from_string(idx.upper()) if isinstance(idx, str) and not idx.isdigit() else int(idx)
    axis = "col" if "col" in fn else "row"
    if fn.startswith("delete"):
        positions = list(range(int(idx), int(idx) + count))
        shift_workbook_formulas(ctx.wb, ws, axis, "del", positions)
    else:
        shift_workbook_formulas(ctx.wb, ws, axis, "ins", [int(idx)], count)
    getattr(ws, fn)(int(idx), count)
    _formula_warn(ctx, ws)
    ctx.note(f"{label} on '{ws.title}' at {p.get('at') or p.get('index')} x{count}")


@op("insert_rows", doc="{sheet, at:row, count?}")
def op_insert_rows(ctx, p):
    _shift(ctx, p, "insert_rows", "Inserted rows")


@op("delete_rows", doc="{sheet, at:row, count?}")
def op_delete_rows(ctx, p):
    _shift(ctx, p, "delete_rows", "Deleted rows")


@op("insert_columns", "insert_cols", doc="{sheet, at:'C', count?}")
def op_insert_cols(ctx, p):
    _shift(ctx, p, "insert_cols", "Inserted columns")


@op("delete_columns", "delete_cols", doc="{sheet, at:'C', count?}")
def op_delete_cols(ctx, p):
    _shift(ctx, p, "delete_cols", "Deleted columns")


@op("named_range", "add_named_range", doc="{name, ref:'Sheet1!$A$1:$A$10'}")
def op_named_range(ctx, p):
    ctx.wb.defined_names[str(p["name"])] = DefinedName(str(p["name"]), attr_text=str(p["ref"]))
    ctx.note(f"Named range '{p['name']}' -> {p['ref']}")


@op("add_comment", "add_note", doc="{sheet, cell, text, author?}")
def op_comment(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    cm = Comment(clean_str(str(p.get("text", ""))), str(p.get("author", "Excel assistant")))
    cm.width, cm.height = 260, 110
    ws[str(p["cell"]).upper()].comment = cm
    ctx.note(f"Comment added to {p['cell']}")


@op("add_hyperlink", "hyperlink", doc="{sheet, cell, url? | targetSheet?+targetCell?, text?}")
def op_hyperlink(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    cell = ws[str(p["cell"]).upper()]
    if p.get("url"):
        cell.hyperlink = str(p["url"])
    else:
        tgt = ctx.sheet(p.get("targetSheet"))
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"{q(tgt.title)}!{p.get('targetCell', 'A1')}")
    if p.get("text"):
        cell.value = str(p["text"])
    cell.font = Font(name=ctx.theme.font, color="0563C1", underline="single")
    ctx.note(f"Hyperlink added at {p['cell']}")


@op("protect_sheet", "protect", doc="{sheet, password?}")
def op_protect(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    ws.protection.sheet = True
    if p.get("password"):
        ws.protection.set_password(str(p["password"]))
    ctx.note(f"Protected '{ws.title}'" + (" with a password" if p.get("password") else ""))


@op("page_setup", "print_setup", doc="{sheet, orientation, fitToWidth, printTitleRows:'1:2', paperSize, footer, header}")
def op_page_setup(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    if p.get("orientation"):
        ws.page_setup.orientation = str(p["orientation"]).lower()
    if p.get("fitToWidth", True):
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    if p.get("printTitleRows"):
        ws.print_title_rows = str(p["printTitleRows"])
    if p.get("paperSize"):
        ws.page_setup.paperSize = {"a4": 9, "letter": 1, "legal": 5, "a3": 8}.get(str(p["paperSize"]).lower(), 9)
    if p.get("footer"):
        ws.oddFooter.center.text = str(p["footer"])
    if p.get("header"):
        ws.oddHeader.center.text = str(p["header"])
    ctx.note(f"Page setup updated on '{ws.title}'")


@op("document_properties", "set_properties", doc="{title, author, subject, keywords}")
def op_doc_props(ctx, p):
    for k in ("title", "subject", "keywords", "description"):
        if p.get(k):
            setattr(ctx.wb.properties, k, str(p[k])[:250])
    if p.get("author"):
        ctx.wb.properties.creator = str(p["author"])[:100]
    ctx.note("Document properties updated")


# ---- rule-based validation ---------------------------------------------------- #
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
URL_RE = re.compile(r"^https?://\S+$", re.I)


@op("validate_data", "validate", doc="{sheet, range?, rules:[{column, required, unique, type:'number'|'integer'|'date'|'text'|'email'|'url', min, max, allowed:[..], regex, maxLength}], highlight?:true}")
def op_validate(ctx, p):
    ws = ctx.sheet(p.get("sheet"))
    reg = op_region(ctx, p, ws)
    rules = p.get("rules") or []
    if not rules:
        raise ExcelError("validate_data needs 'rules'.", status=422, code="INVALID_OPERATION")
    violations: List[Dict[str, Any]] = []
    per_rule: Counter = Counter()
    for rule in rules:
        col = reg.col(rule["column"])
        L = get_column_letter(col)
        header = reg.headers[col - reg.first_col]
        seen: Dict[Any, int] = {}
        for r in range(reg.first_row, reg.last_row + 1):
            v = ws.cell(r, col).value
            blank = v is None or (isinstance(v, str) and not v.strip())
            problems = []
            if blank:
                if rule.get("required"):
                    problems.append("required value is missing")
            elif not is_formula(v):
                t = str(rule.get("type", "")).lower()
                if t in ("number", "integer", "decimal"):
                    num = v if isinstance(v, (int, float)) and not isinstance(v, bool) else None
                    if num is None:
                        problems.append("is not a number")
                    elif t == "integer" and not float(num).is_integer():
                        problems.append("is not a whole number")
                    else:
                        if rule.get("min") is not None and num < rule["min"]:
                            problems.append(f"is below the minimum {rule['min']}")
                        if rule.get("max") is not None and num > rule["max"]:
                            problems.append(f"is above the maximum {rule['max']}")
                elif t == "date":
                    if not isinstance(v, (date, datetime)):
                        problems.append("is not a real date")
                    else:
                        d = v.date() if isinstance(v, datetime) else v
                        if rule.get("min") and d < date.fromisoformat(str(rule["min"])):
                            problems.append(f"is before {rule['min']}")
                        if rule.get("max") and d > date.fromisoformat(str(rule["max"])):
                            problems.append(f"is after {rule['max']}")
                elif t == "email" and not EMAIL_RE.match(str(v).strip()):
                    problems.append("is not a valid email address")
                elif t == "url" and not URL_RE.match(str(v).strip()):
                    problems.append("is not a valid URL")
                if rule.get("allowed") and str(v).strip().lower() not in {str(a).lower() for a in rule["allowed"]}:
                    problems.append("is not an allowed value")
                if rule.get("regex") and not re.fullmatch(str(rule["regex"]), str(v).strip()):
                    problems.append("does not match the required pattern")
                if rule.get("maxLength") and len(str(v)) > int(rule["maxLength"]):
                    problems.append(f"is longer than {rule['maxLength']} characters")
            if rule.get("unique") and not blank:
                key = str(v).strip().lower()
                if key in seen:
                    problems.append(f"duplicates row {seen[key]}")
                else:
                    seen[key] = r
            for msg in problems:
                per_rule[f"{header}: {msg.split(' ')[0] + ' ' + msg.split(' ')[1] if ' ' in msg else msg}"] += 1
                violations.append({"cell": f"{L}{r}", "column": header, "value": None if blank else str(v)[:60], "problem": msg})
    if p.get("highlight", True):
        flagged = {}
        for v in violations[:2000]:
            flagged.setdefault(v["cell"], []).append(v["problem"])
        for i, (coord, msgs) in enumerate(flagged.items()):
            ws[coord].fill = fill_of("FFE699")
            if i < 100:
                ws[coord].comment = Comment("; ".join(msgs)[:250], "Validation")
    ctx.report.setdefault("validation", []).append({
        "sheet": ws.title, "rowsChecked": max(0, reg.last_row - reg.first_row + 1),
        "violations": len(violations), "byRule": dict(per_rule), "examples": violations[:50]})
    ctx.note(f"Validated '{ws.title}': {len(violations)} problem(s) found" + (" (highlighted)" if p.get("highlight", True) and violations else ""))


# ---- automatic dashboard --------------------------------------------------------- #
@op("auto_dashboard", "dashboard", doc="{sheet, range?, name?:'Dashboard', title?} - builds KPI cards, summaries and charts from a data sheet automatically")
def op_dashboard(ctx, p):
    src = ctx.sheet(p.get("sheet"))
    reg = op_region(ctx, p, src)
    if reg.last_row < reg.first_row:
        raise ExcelError("No data rows to build a dashboard from.", status=422, code="EMPTY_SHEET")
    key = "src"
    ctx.tables[key] = reg
    nums = [h for h in numeric_headers(src, reg) if not any(w in h.lower() for w in ("id", "year", "code", "zip", "phone", "no."))]
    cats, date_col = [], None
    for h in reg.headers:
        vals = [src.cell(r, reg.col(h)).value for r in range(reg.first_row, min(reg.last_row, reg.first_row + 2000) + 1)]
        vals = [v for v in vals if v not in (None, "")]
        if not vals:
            continue
        if date_col is None and all(isinstance(v, (date, datetime)) for v in vals):
            date_col = h
        elif all(isinstance(v, str) and not is_formula(v) for v in vals) and 2 <= len({str(v).strip() for v in vals}) <= 12:
            cats.append(h)
    th = ctx.theme
    kpis = [{"label": "Records", "value": f"=COUNTA({key}[{reg.headers[0]}])", "format": "integer"}]
    for h in nums[:3]:
        avg = any(w in h.lower() for w in PERCENT_WORDS + ("score", "rating", "avg", "average", "price", "age", "unit"))
        fmt = reg.formats.get(h) or "number"
        kpis.append({"label": ("Average " if avg else ("" if h.lower().startswith("total") else "Total ")) + h, "value": f"={'AVERAGE' if avg else 'SUM'}({key}[{h}])",
                     "format": fmt if fmt.startswith(('"', "$", "0", "#")) or fmt in ("currency", "currency0", "percent", "integer", "number") else "number"})
    blocks: List[Dict[str, Any]] = [{"type": "title", "text": p.get("title") or f"{src.title} - Dashboard",
                                     "subtitle": f"Generated from '{src.title}' ({reg.last_row - reg.first_row + 1} records)"},
                                    {"type": "kpis", "items": kpis}]
    additive = [h for h in nums if not any(w in h.lower() for w in PERCENT_WORDS + ("score", "rating", "avg", "average", "price", "age", "unit"))]
    values = [{"column": h, "agg": "sum"} for h in (additive or nums)[:2]] or [{"agg": "count", "label": "Records"}]
    for i, cat in enumerate(cats[:2]):
        name = f"dash_{i}"
        blocks.append({"type": "summary", "name": name, "source": key, "groupBy": cat, "values": values,
                       "title": f"{values[0].get('label') or values[0].get('column', 'Records')} by {cat}",
                       "percentOfTotal": bool(nums) and i == 0, "top": 10})
        sname = f"{(values[0].get('agg') or 'sum').title()} of {values[0]['column']}" if values[0].get("column") else "Records"
        blocks.append({"type": "chart", "chartType": "column" if i == 0 else "doughnut" if len(cats) > 1 else "column",
                       "data": {"table": name, "categories": cat, "series": [sname]},
                       "title": f"{sname} by {cat}", "place": "right" if i == 0 else "below"})
    if date_col and nums and reg.last_row - reg.first_row < 400:
        blocks.append({"type": "chart", "chartType": "line", "data": {"table": key, "categories": date_col, "series": [nums[0]]},
                       "title": f"{nums[0]} over time", "place": "below", "width": 24})
    name = safe_sheet_name(p.get("name") or "Dashboard", ctx.wb.sheetnames)
    ws = ctx.wb.create_sheet(name)
    apply_sheet_options(ctx, ws, {"tabColor": th.primary}, 0)
    ws.column_dimensions["A"].width = 2
    render_blocks(ctx, ws, blocks, "B2")
    finalize_tokens(ctx)
    ctx.tables.pop(key, None)
    if p.get("first", True):
        ctx.wb.move_sheet(ws, offset=-(len(ctx.wb.sheetnames) - 1))
        ctx.wb.active = 0
    ctx.note(f"Built dashboard sheet '{name}' with {len(kpis)} KPI(s), {len(cats[:2])} breakdown(s)" + (" and a trend chart" if date_col and nums else ""))


def run_operation(ctx: Ctx, index: int, spec: Any) -> None:
    spec = _maybe_json(spec)
    if not isinstance(spec, dict) or not (spec.get("op") or spec.get("operation") or spec.get("action") or spec.get("type")):
        raise ExcelError(f"Operation {index} must be an object with an 'op' name, e.g. "
                         "{\"op\":\"enhance_sheet\",\"sheet\":\"Sheet1\"}.", status=422, code="INVALID_OPERATION")
    name = str(spec.get("op") or spec.get("operation") or spec.get("action") or spec.get("type"))
    fn = OPS.get(re.sub(r"[^a-z]", "", name.lower()))
    if not fn:
        raise ExcelError(f"Operation {index}: unknown op '{name}'. Available: {', '.join(sorted(OP_DOCS))}.",
                         status=422, code="UNKNOWN_OPERATION")
    try:
        fn(ctx, spec)
    except ExcelError as e:
        e.message = f"Operation {index} ({name}): {e.message}"
        raise
    except (KeyError, ValueError, TypeError, AttributeError) as e:
        raise ExcelError(f"Operation {index} ({name}) has missing or invalid fields ({type(e).__name__}: {str(e)[:120]}). "
                         "Check the field names against /excel-capabilities.", status=422, code="INVALID_OPERATION")


# =========================================================================== #
# PART 3 - analysis / validation report, formula verification, HTTP routes
# =========================================================================== #
ERROR_VALUES = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "#SPILL!", "#CALC!"}
SHEET_REF = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_\.]+))!")
VOLATILE = re.compile(r"\b(OFFSET|INDIRECT|TODAY|NOW|RAND|RANDBETWEEN)\s*\(", re.I)


def jsonable(v: Any) -> Any:
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, float):
        return round(v, 6)
    if isinstance(v, (int, str, bool)) or v is None:
        return v if not isinstance(v, str) else v[:80]
    return str(v)[:80]


def lo_command() -> Optional[List[str]]:
    import shlex
    env = os.environ.get("EXCEL_LO_CMD")
    if env:
        return shlex.split(env)
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    return [exe] if exe else None


def lo_values(data: bytes) -> Optional[Workbook]:
    """Recalculate a COPY with LibreOffice (if installed) and return computed values. None if unavailable."""
    cmd = lo_command()
    if not cmd or os.environ.get("EXCEL_LO_VERIFY", "1") == "0" or len(data) > 5 * 1024 * 1024:
        return None
    try:
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "in.xlsx"
            src.write_bytes(data)
            subprocess.run(cmd + ["--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir",
                                  str(Path(td) / "out"), str(src)], timeout=80, capture_output=True,
                           env={**os.environ, "HOME": td})
            out = Path(td) / "out" / "in.xlsx"
            if not out.exists():
                return None
            import warnings
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                return load_workbook(io.BytesIO(out.read_bytes()), data_only=True)
    except Exception:
        return None


def formula_errors(wb_formulas: Workbook, wb_values: Workbook, limit: int = 60) -> List[Dict[str, str]]:
    found: List[Dict[str, str]] = []
    for ws in wb_formulas.worksheets:
        if ws.title not in wb_values.sheetnames:
            continue
        wv = wb_values[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    v = wv[c.coordinate].value
                    if isinstance(v, str) and v in ERROR_VALUES:
                        found.append({"sheet": ws.title, "cell": c.coordinate, "error": v, "formula": str(c.value)[:100]})
                        if len(found) >= limit:
                            return found
    return found


def analyze_workbook(data: bytes, name: str) -> Dict[str, Any]:
    wb, ext, kind = open_workbook(data, name, data_only=False)
    cached: Optional[Workbook] = None
    if kind == "xlsx":
        try:
            cached = load_workbook(io.BytesIO(data), data_only=True)
        except Exception:
            cached = None
    warnings, feats = package_scan(data) if kind == "xlsx" else ([], {})
    computed = None
    verified = False
    if kind == "xlsx" and (cached is None or any(
            c.data_type == "f" and cached[ws.title][c.coordinate].value is None
            for ws in wb.worksheets[:3] for row in ws.iter_rows(max_row=200) for c in row if c.data_type == "f")):
        computed = lo_values(data)
        verified = computed is not None
    values_wb = computed or cached

    issues: List[Dict[str, Any]] = []

    def add(sev: str, sheet: str, loc: str, typ: str, msg: str) -> None:
        issues.append({"severity": sev, "sheet": sheet, "location": loc, "type": typ, "message": msg})

    sheets_out: List[Dict[str, Any]] = []
    total_formulas = 0
    for ws in wb.worksheets:
        wv = values_wb[ws.title] if values_wb and ws.title in values_wb.sheetnames else None
        info: Dict[str, Any] = {
            "name": ws.title, "state": ws.sheet_state, "dimensions": ws.dimensions, "rows": ws.max_row,
            "columns": ws.max_column, "freezePanes": ws.freeze_panes, "autoFilter": ws.auto_filter.ref or None,
            "mergedRanges": len(ws.merged_cells.ranges), "charts": len(getattr(ws, "_charts", [])),
            "images": len(getattr(ws, "_images", [])), "tables": list(ws.tables.keys()),
            "conditionalFormats": sum(len(r.rules) for r in ws.conditional_formatting),
            "dataValidations": len(ws.data_validations.dataValidation) if ws.data_validations else 0,
            "hiddenColumns": [k for k, d in ws.column_dimensions.items() if d.hidden][:20],
            "hiddenRows": [k for k, d in ws.row_dimensions.items() if d.hidden][:20],
        }
        if ws.sheet_state != "visible":
            add("info", ws.title, "-", "hidden_sheet", f"Sheet '{ws.title}' is {ws.sheet_state}.")
        # ---- formulas
        n_form, volatile, no_cache = 0, 0, 0
        scanned = 0
        for row in ws.iter_rows():
            for c in row:
                scanned += 1
                if scanned > 400_000:
                    break
                if c.data_type != "f":
                    continue
                n_form += 1
                f = c.value.text if hasattr(c.value, "text") else str(c.value)
                if "#REF!" in f:
                    add("error", ws.title, c.coordinate, "broken_reference", f"Formula contains #REF!: {f[:80]}")
                if f.count("(") != f.count(")"):
                    add("error", ws.title, c.coordinate, "unbalanced_parentheses", f"Unbalanced parentheses: {f[:80]}")
                for m in SHEET_REF.finditer(f):
                    ref = (m.group(1) or m.group(2) or "").replace("''", "'")
                    if ref and ref not in wb.sheetnames and not ref.startswith("["):
                        add("error", ws.title, c.coordinate, "missing_sheet", f"Formula refers to sheet '{ref}', which does not exist.")
                        break
                if VOLATILE.search(f):
                    volatile += 1
                if wv is not None:
                    v = wv[c.coordinate].value
                    if isinstance(v, str) and v in ERROR_VALUES:
                        add("error", ws.title, c.coordinate, "formula_error", f"{v} in {c.coordinate}: {f[:80]}")
                    elif v is None:
                        no_cache += 1
        total_formulas += n_form
        info["formulas"] = n_form
        if volatile:
            add("info", ws.title, "-", "volatile_formulas", f"{volatile} volatile formula(s) (OFFSET/INDIRECT/TODAY/NOW/RAND) recalculate constantly and can slow the file.")
        if n_form and wv is None:
            add("info", ws.title, "-", "unverified_formulas", "Formula results could not be checked on the server (no saved values). Open in Excel to confirm.")
        elif n_form and no_cache == n_form and not verified:
            add("info", ws.title, "-", "no_cached_values", "Formulas have no saved results (file written by a script); Excel will calculate on open.")

        # ---- data region profile
        reg = detect_region(ws)
        if reg is None or reg.last_row <= reg.header_row:
            info["region"] = None
            sheets_out.append(info)
            continue
        info["region"] = {"range": reg.data_ref, "headerRow": reg.header_row, "dataRows": reg.last_row - reg.first_row + 1}
        src_ws = wv if (wv is not None and any(True for _ in [0])) else ws
        grid = [list(r) for r in src_ws.iter_rows(min_row=reg.header_row, max_row=reg.last_row, min_col=reg.first_col,
                                                  max_col=reg.last_col, values_only=True)]
        fgrid = [list(r) for r in ws.iter_rows(min_row=reg.header_row, max_row=reg.last_row, min_col=reg.first_col,
                                               max_col=reg.last_col, values_only=True)] if src_ws is not ws else grid
        raw_headers = grid[0]
        for i, h in enumerate(raw_headers):
            if h in (None, ""):
                add("warning", ws.title, f"{get_column_letter(reg.first_col + i)}{reg.header_row}", "blank_header", "A column in the data has no header.")
        low = [str(h).strip().lower() for h in raw_headers if h not in (None, "")]
        for h, k in Counter(low).items():
            if k > 1:
                add("warning", ws.title, f"row {reg.header_row}", "duplicate_header", f"Header '{h}' appears {k} times.")
        body = grid[1:]
        fbody = fgrid[1:]
        blank_rows = [reg.first_row + i for i, row in enumerate(body) if all(v in (None, "") for v in row)]
        if blank_rows:
            add("warning", ws.title, f"row {blank_rows[0]}", "blank_rows", f"{len(blank_rows)} blank row(s) inside the data (first at row {blank_rows[0]}). They break sorting, filters and pivots.")
        keyed = [tuple(str(v).strip().lower() if v is not None else "" for v in row) for row in body if any(v not in (None, "") for v in row)]
        dupes = sum(k - 1 for k in Counter(keyed).values() if k > 1)
        if dupes:
            add("warning", ws.title, reg.data_ref, "duplicate_rows", f"{dupes} fully duplicated row(s).")
        if info["mergedRanges"] and any(reg.first_col <= mr.min_col <= reg.last_col and reg.header_row <= mr.min_row <= reg.last_row for mr in ws.merged_cells.ranges):
            add("warning", ws.title, reg.data_ref, "merged_cells", "Merged cells inside the data block sorting and filtering.")
        cols_out = []
        for ci in range(len(raw_headers)):
            header = str(raw_headers[ci]) if raw_headers[ci] not in (None, "") else f"Column {get_column_letter(reg.first_col + ci)}"
            L = get_column_letter(reg.first_col + ci)
            col = [row[ci] for row in body]
            fcol = [row[ci] for row in fbody]
            nonblank = [v for v in col if v not in (None, "") and not (isinstance(v, str) and not v.strip())]
            n_blank = len(col) - len(nonblank)
            nums = [v for v in nonblank if isinstance(v, (int, float)) and not isinstance(v, bool)]
            dates = [v for v in nonblank if isinstance(v, (date, datetime))]
            strs = [v for v in nonblank if isinstance(v, str)]
            text_nums = [v for v in strs if parse_number(v)]
            text_dates = [v for v in strs if not parse_number(v) and parse_date_text(v)]
            padded = [v for v in strs if v != v.strip() or "  " in v]
            ftypes = Counter(("formula" if is_formula(v) else "value") for v in fcol if v not in (None, ""))
            kinds = [t for t, lst in (("number", nums), ("date", dates), ("text", strs)) if lst]
            ctype = kinds[0] if len(kinds) == 1 else ("mixed" if kinds else "empty")
            entry: Dict[str, Any] = {"header": header, "column": L, "type": ctype, "filled": len(nonblank), "blank": n_blank,
                                     "distinct": len({str(v).strip().lower() for v in nonblank}),
                                     "sample": [jsonable(v) for v in nonblank[:3]]}
            if nums:
                entry.update(min=jsonable(min(nums)), max=jsonable(max(nums)), mean=round(statistics.fmean(nums), 4),
                             sum=jsonable(sum(nums)))
            if dates:
                entry.update(min=jsonable(min(dates)), max=jsonable(max(dates)))
            if ftypes.get("formula"):
                entry["formulas"] = ftypes["formula"]
            cols_out.append(entry)
            if body and n_blank / len(body) > 0.3 and nonblank:
                add("warning" if n_blank / len(body) > 0.6 else "info", ws.title, f"{L}", "many_blanks", f"Column '{header}' is {round(100 * n_blank / len(body))}% empty.")
            if text_nums:
                add("warning", ws.title, f"{L}", "numbers_as_text", f"Column '{header}' has {len(text_nums)} number(s) stored as text (e.g. '{text_nums[0][:20]}'). Sums and charts will ignore them.")
            if text_dates and len(text_dates) >= max(1, len(strs) * 0.5):
                add("warning", ws.title, f"{L}", "dates_as_text", f"Column '{header}' has {len(text_dates)} date(s) stored as text (e.g. '{text_dates[0][:20]}').")
            if padded:
                add("info", ws.title, f"{L}", "extra_spaces", f"Column '{header}' has {len(padded)} value(s) with leading/trailing/double spaces.")
            if len(kinds) > 1 and not (text_nums or text_dates):
                add("warning", ws.title, f"{L}", "mixed_types", f"Column '{header}' mixes {', '.join(kinds)} values.")
            if strs and len(nonblank) >= 8:
                raw_distinct = len({str(v) for v in strs})
                norm_distinct = len({re.sub(r"\s+", " ", v.strip().lower()) for v in strs})
                if norm_distinct < raw_distinct <= 60:
                    variants: Dict[str, set] = {}
                    for v in strs:
                        variants.setdefault(re.sub(r"\s+", " ", v.strip().lower()), set()).add(v)
                    ex = next((sorted(vs) for vs in variants.values() if len(vs) > 1), [])[:2]
                    add("warning", ws.title, f"{L}", "inconsistent_values",
                        f"Column '{header}' has values that differ only by case/spacing (e.g. {' vs '.join(repr(x) for x in ex)}).")
            if len(nums) >= 20:
                qs = statistics.quantiles(nums, n=4)
                iqr = qs[2] - qs[0]
                if iqr > 0:
                    extreme = [v for v in nums if v < qs[0] - 3 * iqr or v > qs[2] + 3 * iqr]
                    if extreme:
                        add("info", ws.title, f"{L}", "outliers", f"Column '{header}' has {len(extreme)} extreme outlier(s) (e.g. {jsonable(extreme[0])}).")
            if ftypes.get("formula", 0) >= 4 and ftypes.get("value", 0) and ftypes["value"] <= 0.2 * (ftypes["formula"] + ftypes["value"]):
                where = [f"{L}{reg.first_row + i}" for i, v in enumerate(fcol) if v not in (None, "") and not is_formula(v)][:5]
                add("warning", ws.title, ", ".join(where), "hardcoded_in_formula_column", f"Column '{header}' is formulas except for {ftypes['value']} typed-in value(s) - likely overwritten formulas.")
        info["columnsProfile"] = cols_out[:60]
        info["sampleRows"] = [[jsonable(v) for v in row] for row in grid[:6]]
        sheets_out.append(info)

    # ---- suggestions (formatting opportunities)
    suggestions: List[str] = []
    for s in sheets_out:
        reg_i = s.get("region")
        if not reg_i:
            continue
        if not s["freezePanes"] and reg_i["dataRows"] > 15:
            suggestions.append(f"'{s['name']}': freeze the header row so it stays visible while scrolling.")
        if not s["autoFilter"] and not s["tables"]:
            suggestions.append(f"'{s['name']}': add filters to the header row.")
        if s["charts"] == 0 and reg_i["dataRows"] >= 3 and any(c["type"] == "number" for c in s.get("columnsProfile", [])):
            suggestions.append(f"'{s['name']}': numeric data could be summarised with a chart or dashboard.")
    if not any(s["conditionalFormats"] for s in sheets_out):
        suggestions.append("No conditional formatting is used; highlighting exceptions (overdue, negative, top/bottom) helps readers.")

    sev_rank = {"error": 0, "warning": 1, "info": 2}
    issues.sort(key=lambda i: sev_rank[i["severity"]])
    counts = Counter(i["severity"] for i in issues)
    kinds_seen = {sev: len({i["type"] for i in issues if i["severity"] == sev}) for sev in sev_rank}
    score = max(0, 100 - min(50, 10 * kinds_seen["error"] + min(5, counts["error"]))
                - min(30, 4 * kinds_seen["warning"]) - min(10, kinds_seen["info"]))
    return {
        "fileName": name, "format": ext.lstrip(".") if kind == "xlsx" else "csv (converted)",
        "sheetCount": len(wb.sheetnames), "totalFormulas": total_formulas, "formulasVerified": verified or bool(cached and total_formulas == 0),
        "healthScore": score,
        "summary": {"errors": counts["error"], "warnings": counts["warning"], "info": counts["info"]},
        "issues": issues[:100], "issuesTruncated": len(issues) > 100,
        "sheets": sheets_out, "suggestions": suggestions[:15], "preservationWarnings": warnings,
        "features": feats,
    }


# --------------------------------------------------------------------------- #
# HTTP routes
# --------------------------------------------------------------------------- #
def slugify(text: Any, default: str = "workbook") -> str:
    return re.sub(r"[^A-Za-z0-9]+", "-", str(text or "")).strip("-")[:50] or default


def create_blueprint(output_dir: Path, base_url_fn: Callable[[], str], logger: Any = None) -> Blueprint:
    import json as _json
    bp = Blueprint("excel", __name__)

    @bp.errorhandler(ExcelError)
    def _err(e: ExcelError):
        if logger:
            try:
                body = request.get_data(as_text=True)
                body = re.sub(r'"fileBase64"\s*:\s*"[^"]{40,}"', '"fileBase64":"<omitted>"', body)
                logger.warning("Excel %s %s -> %s %s | request: %s", request.method, request.path, e.status, e.code, body[:1500])
            except Exception:
                pass
        body = {"status": "error", "errorCode": e.code, "message": e.message}
        if e.details:
            body["details"] = e.details
        return jsonify(body), e.status

    def payload_json() -> Dict[str, Any]:
        data = request.get_json(silent=True)
        if not isinstance(data, dict):
            raise ExcelError("Request body must be a JSON object (Content-Type: application/json).", code="INVALID_JSON")
        if isinstance(data.get("body"), dict) and not any(k in data for k in ("sheets", "operations", "source")):
            data = data["body"]
        return data

    def listish(data: Dict[str, Any], key: str, allow_dict: bool = False) -> Any:
        v = data.get(key)
        if isinstance(v, str):  # Copilot Studio sometimes sends arrays as JSON text
            try:
                v = _json.loads(v)
            except ValueError:
                raise ExcelError(f"'{key}' must be a JSON array.", code="INVALID_JSON")
        if isinstance(v, dict):
            return v if allow_dict else [v]  # a single operation object is treated as a one-item list
        return v if isinstance(v, list) else []

    def cleanup() -> None:
        cutoff = time.time() - int(os.environ.get("FILE_TTL_SECONDS", "86400"))
        for pattern in ("*.xlsx", "*.xlsm"):
            for f in output_dir.glob(pattern):
                try:
                    if f.stat().st_mtime < cutoff:
                        f.unlink()
                except OSError:
                    pass

    def finish(ctx: Ctx, stem: str, ext: str, data: Dict[str, Any], extra: Optional[Dict[str, Any]] = None):
        finalize_tokens(ctx)
        ctx.apply_widths()
        ctx.wb.calculation.fullCalcOnLoad = True
        name = f"{slugify(stem)}-{uuid.uuid4().hex[:8]}{ext}"
        path = output_dir / name
        ctx.wb.save(str(path))
        blob = path.read_bytes()
        verified = False
        if not data.get("skipVerify"):
            computed = lo_values(blob)
            if computed is not None:
                verified = True
                errs = formula_errors(ctx.wb, computed)
                for e in errs[:10]:
                    ctx.warn(f"Formula error {e['error']} at {e['sheet']}!{e['cell']}: {e['formula']}")
                if len(errs) > 10:
                    ctx.warn(f"...and {len(errs) - 10} more formula errors.")
        body: Dict[str, Any] = {
            "status": "success", "fileName": name, "downloadUrl": f"{base_url_fn()}/download/{name}",
            "sheetCount": len(ctx.wb.sheetnames), "sheets": ctx.wb.sheetnames, "changes": ctx.log[:100],
            "warnings": ctx.warnings, "formulasVerified": verified,
        }
        if ctx.report:
            body["report"] = ctx.report
        if extra:
            body.update(extra)
        if data.get("returnBase64") and len(blob) <= 4 * 1024 * 1024:
            body["fileBase64"] = base64.b64encode(blob).decode()
        if logger:
            logger.info("Excel %s (%s sheets)", name, body["sheetCount"])
        return jsonify(body), 200

    @bp.post("/createexcel")
    def createexcel():
        data = payload_json()
        data["sheets"] = listish(data, "sheets", allow_dict=True)
        cleanup()
        ctx = build_workbook(data)
        return finish(ctx, data.get("fileName") or data.get("title") or "workbook", ".xlsx", data)

    @bp.post("/editexcel")
    def editexcel():
        data = payload_json()
        ops = listish(data, "operations")
        if not ops:
            raise ExcelError("'operations' must be a non-empty list. Call /excel-capabilities for the available operations.",
                             status=422, code="NO_OPERATIONS")
        if len(ops) > 80:
            raise ExcelError("Too many operations in one request (max 80).", code="TOO_MANY_OPERATIONS")
        raw, name = load_source(data)
        cleanup()
        wb, ext, kind = open_workbook(raw, name)
        pres_warn, _ = package_scan(raw) if kind == "xlsx" else ([], {})
        ctx = Ctx(wb, Theme(data.get("theme"), data.get("currency")))
        for i, spec in enumerate(ops, 1):
            try:
                run_operation(ctx, i, spec)
            except ExcelError as e:
                if data.get("continueOnError"):
                    ctx.warn(f"Skipped: {e.message}")
                    continue
                raise
        stem = data.get("fileName") or f"{Path(name).stem}-edited"
        return finish(ctx, stem, ext, data, {"preservationWarnings": pres_warn,
                                             "sourceFormat": "csv (converted to xlsx)" if kind == "csv" else ext.lstrip(".")})

    @bp.post("/analyzeexcel")
    def analyzeexcel():
        data = payload_json()
        raw, name = load_source(data)
        return jsonify({"status": "success", **analyze_workbook(raw, name)}), 200

    @bp.get("/excel-capabilities")
    def capabilities():
        return jsonify({
            "status": "success",
            "actions": {"POST /createexcel": "build a workbook from 'sheets'", "POST /editexcel": "apply 'operations' to a source file",
                        "POST /analyzeexcel": "inspect and validate a source file"},
            "blocks": ["title", "text", "kpis", "table", "summary (pivot-style)", "chart", "spacer"],
            "chartTypes": ["column", "bar", "stackedColumn", "stackedBar", "percentStacked", "line", "area", "pie", "doughnut", "scatter", "combo"],
            "conditionalFormats": ["colorScale", "dataBar", "iconSet", "cell", "text", "formula", "topN", "bottomN", "duplicates", "blanks", "status"],
            "validationTypes": ["list", "whole", "decimal", "date", "textLength", "custom"],
            "numberFormats": ["currency", "currency0", "integer", "number", "percent", "percent0", "date", "datetime", "month", "text", "thousands", "millions", "multiple", "accounting"],
            "themes": list(THEMES), "formulaSyntax": {"thisRow": "[@Units]*[@Price]", "tableColumn": "SUM(Sales[Revenue])", "rowNumber": "{r}"},
            "operations": OP_DOCS,
        }), 200

    return bp
